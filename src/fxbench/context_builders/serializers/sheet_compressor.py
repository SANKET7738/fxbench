from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date, datetime, time

from openpyxl.utils import get_column_letter

from fxbench.context_builders.serializers._utils import normalize_value, parse_cell_address
from fxbench.context_builders.serializers.base import IContextBuilder
from fxbench.data_models import Problem

TYPE_LABELS: dict[int, str] = {
    -1: "EmptyCell",
    0: "YearData",
    1: "IntNum",
    2: "FloatNum",
    3: "PercentageNum",
    4: "ScientificNum",
    5: "DateData",
    6: "TimeData",
    7: "CurrencyData",
    8: "EmailData",
    9: "StringData",
}

_RE_YEAR = re.compile(r"^(?:19|20)\d{2}$")
_RE_INT = re.compile(r"^-?\d+$")
_RE_FLOAT = re.compile(r"^-?\d+\.\d+$")
_RE_PERCENTAGE = re.compile(r"^[+-]?(?:\d+(?:\.\d*)?|\.\d+)%$")
_RE_SCIENTIFIC = re.compile(r"^[+-]?(?:\d+(?:\.\d*)?|\.\d+)[eE][+-]?\d+$")
_RE_DATE = re.compile(
    r"^(?:"
    r"\d{4}[-/]\d{1,2}[-/]\d{1,2}"
    r"|\d{1,2}[-/]\d{1,2}[-/]\d{4}"
    r"|\d{1,2}[-/]\d{1,2}"
    r"|\d{4}[-/]\d{1,2}"
    r")$"
)
_RE_TIME = re.compile(r"^(?:2[0-3]|[01]?\d):[0-5]?\d(?:\s?[AaPp][Mm])?$")
_RE_CURRENCY = re.compile(r"^[$\u20ac\u00a3\u00a5]\d+(?:\.\d{1,2})?$")
_RE_EMAIL = re.compile(r"^[\w.\-]+@[\w.\-]+\.\w+$")
_RE_COMMA_NUMBER = re.compile(r"^-?[\d,]+\.?\d*$")
_DIRECTIONS = ((-1, 0), (1, 0), (0, -1), (0, 1))


@dataclass(slots=True)
class _Region:
    type_id: int
    label: str
    min_row: int
    min_col: int
    max_row: int
    max_col: int
    cell_count: int


def _strip_commas(value: str) -> str:
    if "," in value and _RE_COMMA_NUMBER.match(value):
        return value.replace(",", "")
    return value


def classify_cell(value: object) -> tuple[int, str]:
    if value is None:
        return -1, "EmptyCell"
    if isinstance(value, bool):
        return 9, "StringData"
    if isinstance(value, datetime):
        return 5, "DateData"
    if isinstance(value, date):
        return 5, "DateData"
    if isinstance(value, time):
        return 6, "TimeData"
    if isinstance(value, int):
        return (0, "YearData") if 1900 <= value <= 2099 else (1, "IntNum")
    if isinstance(value, float):
        return 2, "FloatNum"

    text = str(value).strip()
    if not text:
        return -1, "EmptyCell"

    number_text = _strip_commas(text)
    if _RE_YEAR.match(number_text):
        return 0, "YearData"
    if _RE_INT.match(number_text):
        return 1, "IntNum"
    if _RE_FLOAT.match(number_text):
        return 2, "FloatNum"
    if _RE_PERCENTAGE.match(text):
        return 3, "PercentageNum"
    if _RE_SCIENTIFIC.match(text):
        return 4, "ScientificNum"
    if _RE_DATE.match(text):
        return 5, "DateData"
    if _RE_TIME.match(text):
        return 6, "TimeData"
    if _RE_CURRENCY.match(text):
        return 7, "CurrencyData"
    if _RE_EMAIL.match(text):
        return 8, "EmailData"
    return 9, "StringData"


def _find_regions(
    type_grid: list[list[int]],
    row_count: int,
    column_count: int,
) -> tuple[list[_Region], list[list[bool]]]:
    visited = [[False] * column_count for _ in range(row_count)]
    in_region = [[False] * column_count for _ in range(row_count)]
    regions: list[_Region] = []

    for row_index in range(row_count):
        for column_index in range(column_count):
            if visited[row_index][column_index]:
                continue

            type_id = type_grid[row_index][column_index]
            if type_id in {-1, 9}:
                visited[row_index][column_index] = True
                continue

            stack = [(row_index, column_index)]
            visited[row_index][column_index] = True
            component: list[tuple[int, int]] = []
            min_row = max_row = row_index
            min_column = max_column = column_index

            while stack:
                current_row, current_column = stack.pop()
                component.append((current_row, current_column))
                min_row = min(min_row, current_row)
                max_row = max(max_row, current_row)
                min_column = min(min_column, current_column)
                max_column = max(max_column, current_column)

                for row_delta, column_delta in _DIRECTIONS:
                    next_row = current_row + row_delta
                    next_column = current_column + column_delta
                    if (
                        0 <= next_row < row_count
                        and 0 <= next_column < column_count
                        and not visited[next_row][next_column]
                        and type_grid[next_row][next_column] == type_id
                    ):
                        visited[next_row][next_column] = True
                        stack.append((next_row, next_column))

            if len(component) >= 2:
                for component_row, component_column in component:
                    in_region[component_row][component_column] = True
                regions.append(
                    _Region(
                        type_id=type_id,
                        label=TYPE_LABELS[type_id],
                        min_row=min_row,
                        min_col=min_column,
                        max_row=max_row,
                        max_col=max_column,
                        cell_count=len(component),
                    )
                )

    return regions, in_region


def _display(value: object) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


class SheetCompressorContextBuilder(IContextBuilder):
    def build(self, problem: Problem) -> str:
        workbook, worksheet = self._open_worksheet(problem)
        try:
            return self._compress(worksheet, problem.active_cell)
        finally:
            workbook.close()

    def _compress(self, worksheet, target_cell: str) -> str:
        target_column, target_row = parse_cell_address(target_cell)
        if target_column is None:
            raise ValueError(f"Invalid cell address: {target_cell}")

        min_row = worksheet.min_row or 1
        max_row = worksheet.max_row or 1
        min_column = worksheet.min_column or 1
        max_column = worksheet.max_column or 1
        row_count = max_row - min_row + 1
        column_count = max_column - min_column + 1

        type_grid: list[list[int]] = [[-1] * column_count for _ in range(row_count)]
        display_grid: list[list[str]] = [[""] * column_count for _ in range(row_count)]
        formula_grid: list[list[bool]] = [[False] * column_count for _ in range(row_count)]

        for row_index, row in enumerate(
            worksheet.iter_rows(
                min_row=min_row,
                max_row=max_row,
                min_col=min_column,
                max_col=max_column,
            )
        ):
            for column_index, cell in enumerate(row):
                row_number = min_row + row_index
                column_number = min_column + column_index
                if row_number == target_row and column_number == target_column:
                    continue

                value = normalize_value(cell)
                is_formula = getattr(cell, "data_type", "n") == "f" or (
                    isinstance(value, str) and value.startswith("=")
                )
                if is_formula:
                    formula_grid[row_index][column_index] = True
                    display_grid[row_index][column_index] = _display(value)
                    continue

                type_id, _ = classify_cell(value)
                type_grid[row_index][column_index] = type_id
                display_grid[row_index][column_index] = _display(value)

        regions, in_region = _find_regions(type_grid, row_count, column_count)

        parts: list[str] = []
        column_letters = [
            get_column_letter(min_column + column_index)
            for column_index in range(column_count)
        ]
        parts.append(f"Spreadsheet has {row_count} rows and {column_count} columns.")
        parts.append(f"Columns: [{', '.join(column_letters)}]")
        if row_count <= 50:
            row_numbers = ", ".join(str(min_row + row_index) for row_index in range(row_count))
        else:
            row_numbers = f"{min_row}, {min_row + 1}, ..., {max_row}"
        parts.append(f"Rows: [{row_numbers}]")
        parts.append("")

        region_tuples = []
        for region in sorted(regions, key=lambda item: (item.min_row, item.min_col)):
            start = f"{get_column_letter(min_column + region.min_col)}{min_row + region.min_row}"
            end = f"{get_column_letter(min_column + region.max_col)}{min_row + region.max_row}"
            region_tuples.append(f"({region.label}|{start}:{end})")

        single_tuples = []
        for row_index in range(row_count):
            for column_index in range(column_count):
                row_number = min_row + row_index
                column_number = min_column + column_index
                address = f"{get_column_letter(column_number)}{row_number}"

                if row_number == target_row and column_number == target_column:
                    single_tuples.append(f"(|{address})")
                elif formula_grid[row_index][column_index]:
                    single_tuples.append(f"({display_grid[row_index][column_index]}|{address})")
                elif not in_region[row_index][column_index] and type_grid[row_index][column_index] != -1:
                    single_tuples.append(f"({display_grid[row_index][column_index]}|{address})")

        if region_tuples:
            parts.append(",".join(region_tuples))
        if single_tuples:
            parts.append(",".join(single_tuples))

        return "\n".join(parts)
