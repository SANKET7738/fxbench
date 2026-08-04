from __future__ import annotations

import datetime
import json

from openpyxl.utils import get_column_letter

from fxbench.context_builders.serializers._utils import normalize_value, parse_cell_address
from fxbench.context_builders.serializers.base import IContextBuilder
from fxbench.data_models import Problem


def _json_default(value):
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return value.isoformat()
    if isinstance(value, datetime.timedelta):
        return str(value)
    return str(value)


class JsonContextBuilder(IContextBuilder):
    def build(self, problem: Problem) -> str:
        workbook, worksheet = self._open_worksheet(problem)
        try:
            return self._build_json(worksheet, problem.active_cell)
        finally:
            workbook.close()

    def _build_json(self, worksheet, target_cell: str) -> str:
        target_column, target_row = parse_cell_address(target_cell)
        if target_column is None:
            raise ValueError(f"Invalid cell address: {target_cell}")

        min_row = worksheet.min_row or 1
        max_row = worksheet.max_row or 1
        min_column = worksheet.min_column or 1
        max_column = worksheet.max_column or 1

        rows = []
        for row_index, row in enumerate(
            worksheet.iter_rows(
                min_row=min_row,
                max_row=max_row,
                min_col=min_column,
                max_col=max_column,
            )
        ):
            row_number = min_row + row_index
            row_data = {}
            for column_index, cell in enumerate(row):
                column_number = min_column + column_index
                cell_id = f"{get_column_letter(column_number)}{row_number}"
                value = normalize_value(cell)
                if row_number == target_row and column_number == target_column:
                    value = ""
                row_data[cell_id] = value
            rows.append(row_data)

        return json.dumps(rows, indent=2, default=_json_default)
