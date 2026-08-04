from __future__ import annotations

from openpyxl.utils import get_column_letter

from fxbench.context_builders.serializers._utils import normalize_value, parse_cell_address
from fxbench.context_builders.serializers.base import IContextBuilder
from fxbench.data_models import Problem


class ViewportContextBuilder(IContextBuilder):
    def __init__(self, row_margin: int = 20, col_margin: int = 10) -> None:
        self.row_margin = row_margin
        self.col_margin = col_margin

    def build(self, problem: Problem) -> str:
        workbook, worksheet = self._open_worksheet(problem)
        try:
            return self._build_viewport(worksheet, problem.active_cell)
        finally:
            workbook.close()

    def _build_viewport(self, worksheet, target_cell: str) -> str:
        target_column, target_row = parse_cell_address(target_cell)
        if target_column is None:
            raise ValueError(f"Invalid cell address: {target_cell}")

        min_row = max(1, target_row - self.row_margin)
        max_row = min(worksheet.max_row or target_row, target_row + self.row_margin)
        min_column = max(1, target_column - self.col_margin)
        max_column = min(worksheet.max_column or target_column, target_column + self.col_margin)

        lines = []
        for row_index, row in enumerate(
            worksheet.iter_rows(
                min_row=min_row,
                max_row=max_row,
                min_col=min_column,
                max_col=max_column,
            )
        ):
            row_number = min_row + row_index
            cells = []
            for column_index, cell in enumerate(row):
                column_number = min_column + column_index
                cell_id = f"{get_column_letter(column_number)}{row_number}"
                value = normalize_value(cell)
                if row_number == target_row and column_number == target_column:
                    value = ""
                cells.append(f"|{cell_id}, {value}")
            lines.append("".join(cells) + "|")

        return "\n".join(lines)
