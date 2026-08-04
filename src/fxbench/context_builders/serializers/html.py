from __future__ import annotations

from openpyxl.utils import get_column_letter

from fxbench.context_builders.serializers._utils import normalize_value, parse_cell_address
from fxbench.context_builders.serializers.base import IContextBuilder
from fxbench.data_models import Problem


class HtmlContextBuilder(IContextBuilder):
    def __init__(
        self,
        min_row: int | None = None,
        max_row: int | None = None,
        min_col: int | None = None,
        max_col: int | None = None,
    ) -> None:
        self._min_row = min_row
        self._max_row = max_row
        self._min_col = min_col
        self._max_col = max_col

    def build(self, problem: Problem) -> str:
        workbook, worksheet = self._open_worksheet(problem)
        try:
            return self._build_html(worksheet, problem.active_cell)
        finally:
            workbook.close()

    def _build_html(self, worksheet, target_cell: str) -> str:
        target_column, target_row = parse_cell_address(target_cell)
        if target_column is None:
            raise ValueError(f"Invalid cell address: {target_cell}")

        min_row = self._min_row or worksheet.min_row or 1
        max_row = self._max_row or worksheet.max_row or 1
        min_column = self._min_col or worksheet.min_column or 1
        max_column = self._max_col or worksheet.max_column or 1

        lines = ["<table>"]
        for row_index, row in enumerate(
            worksheet.iter_rows(
                min_row=min_row,
                max_row=max_row,
                min_col=min_column,
                max_col=max_column,
            )
        ):
            row_number = min_row + row_index
            cells = []
            for column_index, cell in enumerate(row):
                column_number = min_column + column_index
                cell_id = f"{get_column_letter(column_number)}{row_number}"
                value = normalize_value(cell)
                if row_number == target_row and column_number == target_column:
                    cells.append(f'<td id="{cell_id}" class="target"></td>')
                else:
                    cells.append(f'<td id="{cell_id}">{value}</td>')
            lines.append(f'<tr row="{row_number}">{"".join(cells)}</tr>')
        lines.append("</table>")

        return "\n".join(lines)
