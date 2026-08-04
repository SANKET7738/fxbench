from __future__ import annotations

import re
from collections import OrderedDict, defaultdict
from dataclasses import dataclass
from datetime import date, datetime, time
from enum import Enum
from itertools import chain, groupby
from typing import Any

from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter, rows_from_range
from pydantic import BaseModel, ConfigDict, Field

from fxbench.context_builders.serializers._utils import sanitize_formula
from fxbench.context_builders.serializers.base import IContextBuilder
from fxbench.data_models import Problem


@dataclass
class _Cell:
    address: str
    value: object
    format: str | None = None
    formula: str | None = None

    def __hash__(self):
        return hash(self.address)

    def __eq__(self, other):
        if isinstance(other, _Cell):
            return self.address == other.address
        return NotImplemented


class MiniSheetTable(BaseModel):
    model_config = ConfigDict(arbitrary_types_allowed=True)

    address: str
    name: str
    header: list[str]
    data: list[list[_Cell]]


class MiniSheetData(BaseModel):
    model_config = ConfigDict(arbitrary_types_allowed=True)

    values: dict[str, list[str]] = Field(default_factory=dict)
    clusters: dict[str, list[str]] = Field(default_factory=dict)
    formulas: dict[str, list[tuple[str, str]]] = Field(default_factory=dict)
    tables: list[MiniSheetTable] = Field(default_factory=list)


class MiniSheetOrder(str, Enum):
    Row = "row"
    Column = "column"


def _find_blobs(sheet: dict[str, _Cell]) -> list[list[_Cell]]:
    visited: set[str] = set()
    blobs: list[list[_Cell]] = []
    for address, cell in sheet.items():
        if address in visited or cell.value in {None, ""}:
            continue
        blob: list[_Cell] = []
        stack = [address]
        while stack:
            current_address = stack.pop()
            if current_address in visited:
                continue
            visited.add(current_address)
            current_cell = sheet[current_address]
            if current_cell.value is not None:
                blob.append(current_cell)
                for neighbour in _neighbours(current_address):
                    if neighbour in sheet and neighbour not in visited:
                        stack.append(neighbour)
        if blob:
            blobs.append(blob)
    return blobs


def _neighbours(address: str) -> list[str]:
    row, column = coordinate_to_tuple(address)
    neighbours = []
    for delta_column in (-1, 0, 1):
        for delta_row in (-1, 0, 1):
            if delta_column == 0 and delta_row == 0:
                continue
            new_column = column + delta_column
            new_row = row + delta_row
            if new_column >= 1 and new_row >= 1:
                neighbours.append(f"{get_column_letter(new_column)}{new_row}")
    return neighbours


def _bounding(blob: list[_Cell]) -> tuple[int, int, int, int]:
    rows = []
    columns = []
    for cell in blob:
        row, column = coordinate_to_tuple(cell.address)
        rows.append(row)
        columns.append(column)
    return min(rows), max(rows), min(columns), max(columns)


def _edge(
    blob: list[_Cell],
    width_top: int,
    width_bottom: int,
    width_left: int,
    width_right: int,
) -> list[_Cell]:
    if not blob or max(width_top, width_bottom, width_left, width_right) <= 0:
        return []
    min_row, max_row, min_column, max_column = _bounding(blob)
    edge_cells: set[_Cell] = set()
    for cell in blob:
        row, column = coordinate_to_tuple(cell.address)
        if (
            row - min_row < width_top
            or max_row - row < width_bottom
            or column - min_column < width_left
            or max_column - column < width_right
        ):
            edge_cells.add(cell)
    return list(edge_cells)


def _edge_wrapped(
    blob: list[_Cell],
    width_top: int,
    width_bottom: int,
    width_left: int,
    width_right: int,
) -> list[_Cell]:
    if not blob or max(width_top, width_bottom, width_left, width_right) <= 0:
        return []

    mapping = {coordinate_to_tuple(cell.address): cell for cell in blob}
    min_row, max_row, min_column, max_column = _bounding(blob)

    def get_cell(row: int, column: int) -> _Cell | None:
        cell = mapping.get((row, column))
        if cell is not None and cell.value not in {None, ""}:
            return cell
        return None

    def find_cells(row: int, column: int, direction: tuple[int, int], width: int) -> list[_Cell]:
        while min_row <= row <= max_row and min_column <= column <= max_column:
            cell = get_cell(row, column)
            if cell is not None:
                cells = [cell]
                for _ in range(width - 1):
                    row += direction[0]
                    column += direction[1]
                    next_cell = get_cell(row, column)
                    if next_cell is not None:
                        cells.append(next_cell)
                return cells
            row += direction[0]
            column += direction[1]
        return []

    cells: set[_Cell] = set()
    for column in range(min_column, max_column + 1):
        cells.update(find_cells(min_row, column, (1, 0), width_top))
    for column in range(min_column, max_column + 1):
        cells.update(find_cells(max_row, column, (-1, 0), width_bottom))
    for row in range(min_row, max_row + 1):
        cells.update(find_cells(row, min_column, (0, 1), width_left))
    for row in range(min_row, max_row + 1):
        cells.update(find_cells(row, max_column, (0, -1), width_right))
    return list(cells)


def _edge_adaptive(blob: list[_Cell], max_width: int, string_threshold: float) -> list[_Cell]:
    if not blob or max_width <= 0:
        return []

    has_strings = any(isinstance(cell.value, str) and cell.value != "" for cell in blob)
    if not has_strings:
        return _edge_wrapped(blob, 1, 1, 1, 1)

    mapping = {coordinate_to_tuple(cell.address): cell for cell in blob}
    min_row, max_row, min_column, max_column = _bounding(blob)

    def is_string_row(row: int) -> bool:
        row_cells = [mapping[(r, c)] for r, c in mapping if r == row]
        non_empty = [cell for cell in row_cells if cell.value not in {None, ""}]
        if not non_empty:
            return False
        string_count = sum(1 for cell in non_empty if isinstance(cell.value, str))
        return string_count / len(non_empty) >= string_threshold

    def is_string_column(column: int) -> bool:
        column_cells = [mapping[(r, c)] for r, c in mapping if c == column]
        non_empty = [cell for cell in column_cells if cell.value not in {None, ""}]
        if not non_empty:
            return False
        string_count = sum(1 for cell in non_empty if isinstance(cell.value, str))
        return string_count / len(non_empty) >= string_threshold

    width_top = 0
    for row in range(min_row, min(min_row + max_width, max_row + 1)):
        if is_string_row(row):
            width_top += 1
        else:
            break

    width_bottom = 0
    for row in range(max_row, max(max_row - max_width, min_row - 1), -1):
        if is_string_row(row):
            width_bottom += 1
        else:
            break

    width_left = 0
    for column in range(min_column, min(min_column + max_width, max_column + 1)):
        if is_string_column(column):
            width_left += 1
        else:
            break

    width_right = 0
    for column in range(max_column, max(max_column - max_width, min_column - 1), -1):
        if is_string_column(column):
            width_right += 1
        else:
            break

    return _edge_wrapped(
        blob,
        max(width_top, 1),
        max(width_bottom, 1),
        max(width_left, 1),
        max(width_right, 1),
    )


def _format_name(format_string: str) -> str:
    format_lower = format_string.lower()
    if "$" in format_string or "\u20ac" in format_string or "\u00a3" in format_string or "\u00a5" in format_string:
        return "currency"
    if "%" in format_string:
        return "percentage"
    if "e+" in format_lower or "e-" in format_lower:
        return "scientific"
    if "?" in format_string and "/" in format_string:
        return "fraction"
    if any(indicator in format_lower for indicator in ("h:mm", "mm:ss", "[h]", "am/pm")):
        if any(indicator in format_lower for indicator in ("d", "yy", "mmm")):
            return "datetime"
        return "time"
    if any(indicator in format_lower for indicator in ("mm-dd", "dd-mm", "d-mmm", "mmm-yy", "yy", "yyyy", "m/d/")):
        return "date"
    return "number"


def _format_value(value: object, width: int = 32) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        text = value.replace("\n", " ").replace("\r", " ")
        if len(text) > width:
            text = text[: width - 3] + "..."
        return text
    if isinstance(value, (bool, int)):
        return str(value)
    if isinstance(value, float):
        return f"{value:.4f}".rstrip("0").rstrip(".")
    if isinstance(value, (datetime, date, time)):
        return f'"{value.isoformat()}"'
    return repr(value)


def _top_left(cells: list[str]) -> tuple[int, int]:
    min_row = min(coordinate_to_tuple(cell)[0] for cell in cells)
    min_column = min(coordinate_to_tuple(cell)[1] for cell in cells)
    return min_column, min_row


def _compress(cells: list[str]) -> list[str]:
    if not cells:
        return []
    cells = sorted(cells, key=lambda cell: (coordinate_to_tuple(cell)[1], coordinate_to_tuple(cell)[0]))
    compressed = []
    for _, group in groupby(cells, key=lambda cell: coordinate_to_tuple(cell)[1]):
        group = list(group)
        current_sequence = [group[0]]
        sequences = []
        for index in range(1, len(group)):
            previous_row = coordinate_to_tuple(group[index - 1])[0]
            current_row = coordinate_to_tuple(group[index])[0]
            if current_row == previous_row + 1:
                current_sequence.append(group[index])
            else:
                sequences.append(current_sequence)
                current_sequence = [group[index]]
        sequences.append(current_sequence)
        for sequence in sequences:
            if len(sequence) == 1:
                compressed.append(sequence[0])
            else:
                compressed.append(f"{sequence[0]}:{sequence[-1]}")
    return compressed


_CELL_REF_RE = re.compile(
    r"(\$?[A-Z]{1,3})\$?(\d+)"
    r"(?::(\$?[A-Z]{1,3})\$?(\d+))?",
    re.IGNORECASE,
)


def _get_column(reference: str) -> str:
    return reference.lstrip("$").upper()


def _group_formula(formula: str) -> str:
    def replace_reference(match: re.Match) -> str:
        first_column = _get_column(match.group(1))
        if match.group(3):
            second_column = _get_column(match.group(3))
            return f"{first_column}@:{second_column}@"
        return f"{first_column}@"

    return _CELL_REF_RE.sub(replace_reference, formula)


def _extract_cells(worksheet, workbook_path: str | None = None) -> dict[str, _Cell]:
    from openpyxl import load_workbook
    from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

    values_workbook = None
    values_worksheet = None
    if workbook_path:
        try:
            values_workbook = load_workbook(workbook_path, data_only=True, read_only=False)
            values_worksheet = values_workbook[worksheet.title]
        except Exception:
            values_workbook = None
            values_worksheet = None

    cells: dict[str, _Cell] = {}
    try:
        source_cells = getattr(worksheet, "_cells", None)
        if source_cells is not None:
            iterable = sorted(source_cells.values(), key=lambda cell: (cell.row, cell.column))
        else:
            iterable = (
                cell
                for row in worksheet.iter_rows(
                    min_row=worksheet.min_row,
                    max_row=worksheet.max_row,
                    min_col=worksheet.min_column,
                    max_col=worksheet.max_column,
                )
                for cell in row
            )

        value_cells = getattr(values_worksheet, "_cells", {}) if values_worksheet is not None else {}
        for cell in iterable:
            address = f"{get_column_letter(cell.column)}{cell.row}"
            formula = None
            if cell.data_type == "f":
                raw = cell.value
                if isinstance(raw, ArrayFormula):
                    formula = raw.text
                elif isinstance(raw, DataTableFormula):
                    formula = None
                elif raw is not None:
                    formula = str(raw)
                if formula and not formula.startswith("="):
                    formula = f"={formula}"
                if formula:
                    formula = sanitize_formula(formula)

            if values_worksheet is not None:
                evaluated_cell = value_cells.get((cell.row, cell.column))
                value = evaluated_cell.value if evaluated_cell is not None else None
            else:
                value = cell.value if formula is None else None

            if value in {None, ""} and formula is None:
                continue

            number_format = cell.number_format
            if number_format == "General":
                number_format = None
            elif number_format and number_format.startswith('"$"'):
                number_format = number_format.replace('"$"', "$")

            cells[address] = _Cell(
                address=address,
                value=value,
                format=number_format,
                formula=formula,
            )
    finally:
        if values_workbook is not None:
            values_workbook.close()

    return cells


def _validate_cell_address(address: str) -> None:
    try:
        coordinate_to_tuple(address)
    except Exception:
        raise ValueError(f"Invalid cell address: {address}") from None


class MiniSheetContextBuilder(IContextBuilder):
    _read_only: bool = False

    def __init__(
        self,
        width: int = 2,
        wrap: bool = True,
        order: MiniSheetOrder = MiniSheetOrder.Row,
        delimiter: str = " | ",
        inverse: bool = False,
        numbers: bool = False,
        width_top: int | None = None,
        width_bottom: int | None = None,
        width_left: int | None = None,
        width_right: int | None = None,
        adaptive: bool = False,
        max_width: int = 5,
        string_threshold: float = 0.75,
        include_strings: bool = True,
        include_clusters: bool = True,
        include_formulas: bool = True,
        include_tables: bool = True,
        cache_size: int = 8,
    ) -> None:
        self.width = width
        self.wrap = wrap
        self.order = order
        self.delimiter = delimiter
        self.inverse = inverse
        self.numbers = numbers
        self.width_top = width_top if width_top is not None else width
        self.width_bottom = width_bottom if width_bottom is not None else width
        self.width_left = width_left if width_left is not None else width
        self.width_right = width_right if width_right is not None else width
        self.adaptive = adaptive
        self.max_width = max_width
        self.string_threshold = string_threshold
        self.include_strings = include_strings
        self.include_clusters = include_clusters
        self.include_formulas = include_formulas
        self.include_tables = include_tables
        self.cache_size = cache_size
        self._sheet_cache: OrderedDict[tuple[str, str], tuple[Any, Any, dict[str, _Cell]]] = OrderedDict()

    def build(self, problem: Problem) -> str:
        _validate_cell_address(problem.active_cell)
        workbook, worksheet, base_sheet_cells = self._get_sheet_data(problem)
        try:
            sheet_cells = base_sheet_cells.copy()
            if problem.active_cell in sheet_cells:
                sheet_cells[problem.active_cell] = _Cell(
                    address=problem.active_cell,
                    value=None,
                    formula=None,
                    format=None,
                )
            data = self._collect(sheet_cells, worksheet, problem.worksheet_name)
            return self._compile(data)
        finally:
            if self.cache_size <= 0:
                workbook.close()

    def close(self) -> None:
        while self._sheet_cache:
            _, (workbook, _, _) = self._sheet_cache.popitem(last=False)
            workbook.close()

    def _get_sheet_data(self, problem: Problem) -> tuple[Any, Any, dict[str, _Cell]]:
        key = (str(problem.workbook_path), problem.worksheet_name)
        cached = self._sheet_cache.get(key)
        if cached is not None:
            self._sheet_cache.move_to_end(key)
            return cached

        workbook, worksheet = self._open_worksheet(problem)
        try:
            sheet_cells = _extract_cells(worksheet, workbook_path=str(problem.workbook_path))
        except Exception:
            workbook.close()
            raise

        if self.cache_size <= 0:
            return workbook, worksheet, sheet_cells

        self._sheet_cache[key] = (workbook, worksheet, sheet_cells)
        self._sheet_cache.move_to_end(key)
        while len(self._sheet_cache) > self.cache_size:
            _, (old_workbook, _, _) = self._sheet_cache.popitem(last=False)
            old_workbook.close()
        return workbook, worksheet, sheet_cells

    def _collect(
        self,
        sheet_cells: dict[str, _Cell],
        worksheet,
        sheet_name: str,
    ) -> MiniSheetData:
        del sheet_name

        empty: set[str] = set()
        exclude: set[str] = set()
        tables: list[MiniSheetTable] = []

        for table in worksheet.tables.values():
            table_rows = list(rows_from_range(table.ref))
            header_row = table_rows[0]
            data_rows = table_rows[1:2]
            tables.append(
                MiniSheetTable(
                    address=table.ref,
                    name=table.name,
                    header=[
                        str(sheet_cells[cell].value)
                        if cell in sheet_cells and sheet_cells[cell].value is not None
                        else ""
                        for cell in header_row
                    ],
                    data=[
                        [
                            sheet_cells.get(cell, _Cell(address=cell, value=None))
                            for cell in row
                        ]
                        for row in data_rows
                    ],
                )
            )
            exclude.update(cell for row in table_rows for cell in row)

        blobs = _find_blobs(sheet_cells)
        if self.adaptive:
            blob_cells = chain.from_iterable(
                _edge_adaptive(blob, self.max_width, self.string_threshold)
                for blob in blobs
            )
        else:
            edge_function = _edge_wrapped if self.wrap else _edge
            blob_cells = chain.from_iterable(
                edge_function(
                    blob,
                    self.width_top,
                    self.width_bottom,
                    self.width_left,
                    self.width_right,
                )
                for blob in blobs
            )

        values: dict[str, list[str]] = defaultdict(list)
        for cell in blob_cells:
            if cell.value in {None, ""} or cell.address in exclude:
                continue
            if not isinstance(cell.value, str) and not self.numbers:
                continue
            values[_format_value(cell.value)].append(cell.address)

        clusters: dict[str, list[str]] = defaultdict(list)
        for cell in sheet_cells.values():
            if cell.value in {None, ""}:
                empty.add(cell.address)
                continue
            if isinstance(cell.value, str) or cell.address in exclude:
                continue
            clusters[_format_name(cell.format or "")].append(cell.address)

        formulas: dict[str, list[tuple[str, str]]] = defaultdict(list)
        for cell in sheet_cells.values():
            if cell.formula is None or cell.address in exclude:
                continue
            formulas[_group_formula(cell.formula)].append((cell.address, cell.formula))

        return MiniSheetData(
            values=values,
            clusters=clusters,
            formulas=formulas,
            tables=tables,
        )

    def _compile(self, components: MiniSheetData) -> str:
        compiled = ""
        if self.include_strings and components.values:
            compiled += "The following (string) values were found on the edges of data clusters (grouped by unique value, truncated if applicable).\n\n"
            for value, cells in sorted(components.values.items(), key=lambda item: (self._sorter(item[1]), str(item[0]))):
                compressed_cells = ",".join(_compress(cells))
                left, right = value, compressed_cells
                if not self.inverse:
                    left, right = right, left
                compiled += f"{left}{self.delimiter}{right}\n"

        if self.include_clusters and components.clusters:
            compiled += "\nThe following (ranges of) numerical values were found (grouped by their format type).\n\n"
            for format_name, cells in sorted(components.clusters.items(), key=lambda item: (self._sorter(item[1]), str(item[0]))):
                compressed_cells = ",".join(_compress(cells))
                left, right = format_name, compressed_cells
                if not self.inverse:
                    left, right = right, left
                compiled += f"{left}{self.delimiter}{right}\n"

        if self.include_formulas and components.formulas:
            compiled += "\nThe following formulas were found (grouped by their formula pattern if applicable).\n\n"
            for pattern, entries in sorted(
                components.formulas.items(),
                key=lambda item: (self._sorter([address for address, _ in item[1]]), str(item[0])),
            ):
                formula_text = entries[0][1] if len(entries) == 1 else pattern
                compressed_cells = ",".join(_compress([address for address, _ in entries]))
                left, right = formula_text, compressed_cells
                if not self.inverse:
                    left, right = right, left
                compiled += f"{left}{self.delimiter}{right}\n"

        if self.include_tables and components.tables:
            compiled += "\nThe following tables (and one sample row) were found in the workbook.\n\n"
            for table in components.tables:
                compiled += f"- Table '{table.name}' at {table.address}.\n"
                compiled += f"  {self.delimiter.join(str(header) for header in table.header)}\n"
                if table.data:
                    compiled += f"  {self.delimiter.join(_format_value(cell.value) for cell in table.data[0])}\n"

        return compiled.strip()

    def _sorter(self, cells: list[str]) -> tuple[int, int]:
        column, row = _top_left(cells)
        if self.order == MiniSheetOrder.Row:
            return row, column
        return column, row


__all__ = [
    "MiniSheetContextBuilder",
    "MiniSheetOrder",
]
