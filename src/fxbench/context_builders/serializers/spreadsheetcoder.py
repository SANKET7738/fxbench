from __future__ import annotations

from openpyxl.utils import get_column_letter, rows_from_range

from fxbench.context_builders.serializers._utils import normalize_value, parse_cell_address
from fxbench.context_builders.serializers.base import IContextBuilder
from fxbench.data_models import Problem


class SpreadsheetCoderContextBuilder(IContextBuilder):
    _read_only: bool = False

    ROW_MARGIN = 10
    COL_MARGIN = 10

    def build(self, problem: Problem) -> str:
        workbook, worksheet = self._open_worksheet(problem)
        try:
            return self._build_viewport(worksheet, problem.active_cell)
        finally:
            workbook.close()

    def _build_viewport(self, worksheet, target_cell: str) -> str:
        target_column, target_row = parse_cell_address(target_cell)
        if target_column is None:
            raise ValueError(f"Invalid cell address: {target_cell}")

        min_row = max(1, target_row - self.ROW_MARGIN)
        max_row = min(worksheet.max_row or target_row, target_row + self.ROW_MARGIN)
        min_column = max(1, target_column - self.COL_MARGIN)
        max_column = min(worksheet.max_column or target_column, target_column + self.COL_MARGIN)

        header_bounds = self._find_header_bounds(worksheet, target_column, target_row)
        header_row = header_bounds[0] if header_bounds is not None else None
        header_min_column = header_bounds[1] if header_bounds is not None else None
        header_max_column = header_bounds[2] if header_bounds is not None else None
        header_needs_section = (
            header_bounds is not None
            and (
                header_row < min_row
                or header_row > max_row
                or header_min_column < min_column
                or header_max_column > max_column
            )
        )

        sections = []
        if header_needs_section:
            sections.append("## Column Headers")
            sections.append(
                self._render_row(
                    worksheet,
                    header_row,
                    header_min_column,
                    header_max_column,
                    target_column,
                    target_row,
                )
            )
            sections.append("")

        sections.append("## Data")
        for row_number in range(min_row, max_row + 1):
            sections.append(
                self._render_row(
                    worksheet,
                    row_number,
                    min_column,
                    max_column,
                    target_column,
                    target_row,
                )
            )

        return "\n".join(sections)

    def _render_row(
        self,
        worksheet,
        row_number: int,
        min_column: int,
        max_column: int,
        target_column: int,
        target_row: int,
    ) -> str:
        cells = []
        for column_number in range(min_column, max_column + 1):
            cell_id = f"{get_column_letter(column_number)}{row_number}"
            value = normalize_value(worksheet.cell(row=row_number, column=column_number))
            if row_number == target_row and column_number == target_column:
                value = ""
            cells.append(f"|{cell_id}, {value}")

        return "".join(cells) + "|"

    def _find_header_bounds(
        self,
        worksheet,
        target_column: int,
        target_row: int,
    ) -> tuple[int, int, int] | None:
        for table in worksheet.tables.values():
            table_rows = list(rows_from_range(table.ref))
            if not table_rows:
                continue

            first_cell = table_rows[0][0]
            last_cell = table_rows[-1][-1]
            min_column, min_row = parse_cell_address(first_cell)
            max_column, max_row = parse_cell_address(last_cell)

            if (
                min_column is not None
                and max_column is not None
                and min_column <= target_column <= max_column
                and min_row <= target_row <= max_row
            ):
                return min_row, min_column, max_column

        return None

    def _find_header_row(self, worksheet, target_column: int, target_row: int) -> int | None:
        bounds = self._find_header_bounds(worksheet, target_column, target_row)
        return bounds[0] if bounds is not None else None
