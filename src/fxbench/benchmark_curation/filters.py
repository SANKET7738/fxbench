from __future__ import annotations

import re
import zipfile
from collections import defaultdict
from io import BytesIO
from typing import Dict, List, Tuple

import openpyxl
from pycel.lib.function_info import all_excel_functions

from fxbench.utils import get_formula_text, is_predominantly_english


VALID_EXCEL_FUNCTIONS: frozenset[str] = frozenset(f.upper() for f in all_excel_functions)

_FUNC_RE = re.compile(r"([A-Za-z][A-Za-z0-9_.]*)\s*\(")

def file_level_filters(
    xlsx_bytes: bytes,
    max_size_mb: int = 50,
    max_sheets: int = 50,
) -> Tuple[bool, List[str], Dict]:
    reasons: List[str] = []
    metadata: Dict = {
        "file_size_kb": len(xlsx_bytes) / 1024,
        "is_valid_zip": False,
        "has_encryption": False,
        "has_vba": False,
        "sheet_count": None,
    }

    size_mb = len(xlsx_bytes) / (1024 * 1024)
    if size_mb > max_size_mb:
        reasons.append(f"too large: {size_mb:.2f} MB")
    if size_mb < 0.001:
        reasons.append(f"too small: {size_mb:.2f} MB")

    buf = BytesIO(xlsx_bytes)
    if not zipfile.is_zipfile(buf):
        reasons.append("not a valid ZIP archive")
        return False, reasons, metadata
    metadata["is_valid_zip"] = True

    buf.seek(0)
    try:
        with zipfile.ZipFile(buf, "r") as zf:
            namelist = zf.namelist()
            has_content_types = "[Content_Types].xml" in namelist
            has_xl_dir = any(n.startswith("xl/") for n in namelist)
            if not has_content_types or not has_xl_dir:
                if "EncryptedPackage" in namelist or "EncryptionInfo" in namelist:
                    reasons.append("file is encrypted")
                    metadata["has_encryption"] = True
                else:
                    reasons.append("invalid xlsx structure")
            if any("vbaProject" in n for n in namelist):
                metadata["has_vba"] = True
                reasons.append("has_vba_macros")
    except zipfile.BadZipFile:
        reasons.append("corrupted ZIP archive")
        return False, reasons, metadata
    except Exception as e:
        reasons.append(f"zip error: {e}")
        metadata["has_encryption"] = True
        return False, reasons, metadata

    if not reasons:
        buf.seek(0)
        try:
            wb = openpyxl.load_workbook(buf, read_only=True)
            metadata["sheet_count"] = len(wb.sheetnames)
            wb.close()
            if metadata["sheet_count"] > max_sheets:
                reasons.append(f"too many sheets: {metadata['sheet_count']}")
            if metadata["sheet_count"] == 0:
                reasons.append("no sheets")
        except Exception as e:
            reasons.append(f"openpyxl load error: {e}")

    return len(reasons) == 0, reasons, metadata

def is_sheet_valid(
    ws,
    min_cell_count: int = 10,
    max_cell_count: int = 50_000,
    min_formula_ratio: float = 0.03,
    max_formula_ratio: float = 0.9,
) -> Tuple[bool, List[str], Dict]:
    non_empty = 0
    formulas = 0
    string_cells: List[str] = []
    has_external_links = False

    try:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                non_empty += 1
                if cell.data_type == "f":
                    formulas += 1
                    if not has_external_links:
                        s = str(cell.value)
                        if "[" in s or "http" in s.lower():
                            has_external_links = True
                elif isinstance(cell.value, str) and len(cell.value.strip()) >= 2:
                    string_cells.append(cell.value.strip())
    except Exception:
        return False, ["error iterating rows"], {}

    formula_ratio = formulas / non_empty if non_empty else 0.0
    metadata = {
        "non_empty_cells": non_empty,
        "formula_cells": formulas,
        "formula_ratio": round(formula_ratio, 4),
        "has_external_links": has_external_links,
    }

    if non_empty < min_cell_count:
        return False, ["too small"], metadata
    if non_empty > max_cell_count:
        return False, ["too large"], metadata
    if formulas < 2:
        return False, ["too few formulas"], metadata
    if formula_ratio > max_formula_ratio:
        return False, ["template"], metadata
    if formula_ratio < min_formula_ratio:
        return False, ["data dump"], metadata
    if has_external_links:
        return False, ["has_external_links"], metadata
    if not is_predominantly_english(string_cells):
        return False, ["non_english_content"], metadata

    return True, [], metadata


def sheet_level_filters(xlsx_bytes: bytes) -> List[Dict]:
    valid: List[Dict] = []
    try:
        wb = openpyxl.load_workbook(BytesIO(xlsx_bytes), read_only=True)
        for name in wb.sheetnames:
            ok, _, metadata = is_sheet_valid(wb[name])
            if ok:
                valid.append({"sheet_name": name, "metadata": metadata})
        wb.close()
    except Exception:
        pass
    return valid


def extract_functions(formula: str) -> set[str]:
    return {m.upper() for m in _FUNC_RE.findall(formula)}


def is_sheet_formulas_valid(ws) -> Tuple[bool, Dict]:
    all_funcs: dict[str, int] = defaultdict(int)
    invalid_funcs: dict[str, int] = defaultdict(int)
    total = 0
    valid = True

    try:
        for row in ws.iter_rows():
            for cell in row:
                text = get_formula_text(cell)
                if text is None:
                    continue
                total += 1
                for fn in extract_functions(text):
                    all_funcs[fn] += 1
                    if fn not in VALID_EXCEL_FUNCTIONS:
                        invalid_funcs[fn] += 1
                        valid = False
    except Exception:
        return False, {
            "all_functions": dict(all_funcs),
            "invalid_functions": dict(invalid_funcs),
            "total_formulas": total,
            "error": "iteration_error",
        }

    return valid, {
        "all_functions": dict(all_funcs),
        "invalid_functions": dict(invalid_funcs),
        "total_formulas": total,
    }
