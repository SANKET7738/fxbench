from __future__ import annotations

import re
from collections import Counter, defaultdict
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

import openpyxl
from openpyxl.utils.cell import coordinate_to_tuple, rows_from_range
from openpyxl.utils.exceptions import CellCoordinatesException, InvalidFileException


MAX_PER_SKETCH = 5
GLOBAL_MEDIAN_SHEET_SIZE = 182

_RE_STRUCTURED = re.compile(r"(?:[\w.]+\[)|\[@", re.VERBOSE)
_RE_CROSS_SHEET = re.compile(r"(?:'[^']+' | [\w.]+)!", re.VERBOSE)
_RE_WHOLE_ROW_COL = re.compile(
    r"(?<![A-Z0-9])(?:\$?[A-Z]{1,3}:\$?[A-Z]{1,3}|\$?\d+:\$?\d+)(?![A-Z0-9])",
    re.VERBOSE,
)
_RE_ABSOLUTE = re.compile(r"\$[A-Z]{1,3}\$?\d+|[A-Z]{1,3}\$\d+", re.VERBOSE)
_RE_RANGE = re.compile(r"(?<![A-Z])\$?[A-Z]{1,3}\$?\d+:\$?[A-Z]{1,3}\$?\d+", re.VERBOSE)
_RE_CELL = re.compile(r"(?<![A-Z$:])[A-Z]{1,3}\d+(?![\d:$])", re.VERBOSE)
_CELL_PATTERN = re.compile(r"^[A-Z]{1,3}\d+$")
_COL_PATTERN = re.compile(r"^[A-Z]{1,3}$")
_RE_NAMED_IDENT = re.compile(r"(?<![A-Z0-9_$])([A-Z_][A-Z0-9_.]+)(?![A-Z0-9_.])")
_RE_STRING_LIT = re.compile(r'"[^"]*"')
_RE_STRUCT_FULL = re.compile(r"[\w.]+\[[^\]]*\]|\[@[^\]]*\]")


def neighbor_cells(row: int, column: int) -> list[tuple[int, int]]:
    cells = []
    for row_delta in (-1, 0, 1):
        for column_delta in (-1, 0, 1):
            if row_delta == 0 and column_delta == 0:
                continue
            neighbor_row = row + row_delta
            neighbor_column = column + column_delta
            if neighbor_row >= 1 and neighbor_column >= 1:
                cells.append((neighbor_row, neighbor_column))
    return cells


def find_blobs(filled_cells: set[tuple[int, int]]) -> list[set[tuple[int, int]]]:
    visited = set()
    blobs = []
    for position in filled_cells:
        if position in visited:
            continue
        blob = set()
        stack = [position]
        while stack:
            current = stack.pop()
            if current in visited:
                continue
            visited.add(current)
            blob.add(current)
            for neighbor in neighbor_cells(current[0], current[1]):
                if neighbor in filled_cells and neighbor not in visited:
                    stack.append(neighbor)
        if blob:
            blobs.append(blob)
    return blobs


def cells_in_range(range_text: str) -> set[tuple[int, int]]:
    cells = set()
    try:
        row_groups = rows_from_range(range_text)
    except (ValueError, CellCoordinatesException):
        return cells
    for row_group in row_groups:
        for address in row_group:
            cells.add(coordinate_to_tuple(address))
    return cells


def has_cell_within_radius(
    active_cell: tuple[int, int],
    cell_group: set[tuple[int, int]],
    radius: int = 5,
) -> bool:
    active_row, active_column = active_cell
    for row, column in cell_group:
        if abs(active_row - row) <= radius and abs(active_column - column) <= radius:
            return True
    return False


def classify_location(
    worksheet,
    active_row: int,
    active_column: int,
    filled_cells: set[tuple[int, int]],
    blobs: list[set[tuple[int, int]]],
) -> str:
    active_cell = (active_row, active_column)
    table_cells = set()
    for table in worksheet.tables.values():
        table_cells_for_ref = cells_in_range(table.ref)
        if active_cell in table_cells_for_ref:
            return "ctrl_t_table"
        table_cells.update(table_cells_for_ref)
    if table_cells and has_cell_within_radius(active_cell, table_cells):
        return "ctrl_t_adjacent"
    for blob in blobs:
        if active_cell in blob:
            return "blob_inside"
    all_blob_cells = set()
    for blob in blobs:
        all_blob_cells.update(blob)
    if all_blob_cells and has_cell_within_radius(active_cell, all_blob_cells):
        return "blob_adjacent"
    return "unstructured"


def classify_formula_refs(formula: str) -> list[str]:
    if not formula:
        return []
    formula_without_equals = formula.lstrip("=") if formula.startswith("=") else formula
    formula_upper = formula_without_equals.upper()
    ref_types = []

    if _RE_STRUCTURED.search(formula_without_equals):
        ref_types.append("structured_ref")
    if _RE_CROSS_SHEET.search(formula_without_equals):
        ref_types.append("cross_sheet_ref")
    if _RE_WHOLE_ROW_COL.search(formula_upper):
        ref_types.append("whole_row_col_ref")
    if _RE_ABSOLUTE.search(formula_upper):
        ref_types.append("absolute_ref")

    formula_without_whole_row_col = _RE_WHOLE_ROW_COL.sub("", formula_upper)
    if _RE_RANGE.search(formula_without_whole_row_col):
        ref_types.append("range_ref")

    formula_without_refs = _RE_RANGE.sub("", formula_without_whole_row_col)
    formula_without_refs = _RE_CROSS_SHEET.sub("", formula_without_refs)
    formula_without_refs = _RE_STRUCTURED.sub("", formula_without_refs)
    formula_without_refs = _RE_ABSOLUTE.sub("", formula_without_refs)
    if _RE_CELL.search(formula_without_refs):
        ref_types.append("cell_ref")

    formula_without_literals = _RE_STRING_LIT.sub("", formula_upper)
    formula_without_literals = _RE_STRUCT_FULL.sub("", formula_without_literals)
    formula_without_literals = _RE_CROSS_SHEET.sub("", formula_without_literals)
    for match in _RE_NAMED_IDENT.finditer(formula_without_literals):
        token = match.group(1)
        token_rest = formula_without_literals[match.end() :].lstrip()
        if token_rest.startswith("("):
            continue
        if _CELL_PATTERN.match(token) or _COL_PATTERN.match(token):
            continue
        ref_types.append("named_ref")
        break

    return ref_types if ref_types else ["cell_ref"]


def filled_cells_for_sheet(worksheet) -> set[tuple[int, int]]:
    filled_cells = set()
    for row in worksheet.iter_rows():
        for cell in row:
            if cell.value is not None:
                filled_cells.add((cell.row, cell.column))
    return filled_cells


def fallback_annotations(
    cells: list[dict],
    sheet_size: int,
    fallback_reason: str,
) -> list[dict]:
    results = []
    for cell in cells:
        results.append(
            {
                "idx": cell["idx"],
                "location": "unstructured",
                "formula_ref_types": classify_formula_refs(cell["formula"]),
                "sheet_size": sheet_size,
                "fallback_reason": fallback_reason,
            }
        )
    return results


def fallback_annotations_for_sheet_items(
    sheet_items: list[tuple[str, list[dict]]],
    fallback_reason: str,
) -> list[dict]:
    results = []
    for _, cells in sheet_items:
        results.extend(fallback_annotations(cells, 0, fallback_reason))
    return results


def process_workbook(workbook_path: str, sheet_items: list[tuple[str, list[dict]]]) -> list[dict]:
    if not Path(workbook_path).exists():
        return fallback_annotations_for_sheet_items(sheet_items, "missing_workbook")

    workbook = None
    try:
        workbook = openpyxl.load_workbook(workbook_path, read_only=False, data_only=False)
    except (BadZipFile, InvalidFileException, KeyError, OSError, ParseError, ValueError):
        return fallback_annotations_for_sheet_items(sheet_items, "workbook_load_error")

    try:
        sheet_cache = {}
        results = []
        for sheet_name, cells in sheet_items:
            if sheet_name not in workbook.sheetnames:
                results.extend(fallback_annotations(cells, 0, "missing_sheet"))
                continue

            if sheet_name not in sheet_cache:
                worksheet = workbook[sheet_name]
                filled_cells = filled_cells_for_sheet(worksheet)
                sheet_cache[sheet_name] = {
                    "worksheet": worksheet,
                    "filled_cells": filled_cells,
                    "blobs": find_blobs(filled_cells),
                    "sheet_size": len(filled_cells),
                }

            cached_sheet = sheet_cache[sheet_name]
            for cell in cells:
                try:
                    active_row, active_column = coordinate_to_tuple(cell["cell"])
                except (ValueError, CellCoordinatesException):
                    results.extend(
                        fallback_annotations(
                            [cell],
                            cached_sheet["sheet_size"],
                            "invalid_cell",
                        )
                    )
                    continue

                location = classify_location(
                    cached_sheet["worksheet"],
                    active_row,
                    active_column,
                    cached_sheet["filled_cells"],
                    cached_sheet["blobs"],
                )
                results.append(
                    {
                        "idx": cell["idx"],
                        "location": location,
                        "formula_ref_types": classify_formula_refs(cell["formula"]),
                        "sheet_size": cached_sheet["sheet_size"],
                        "fallback_reason": None,
                    }
                )
        return results
    finally:
        workbook.close()


def size_bucket(sheet_size: int) -> str:
    if sheet_size < 50:
        return "XS"
    if sheet_size <= 200:
        return "S"
    if sheet_size <= 500:
        return "M"
    if sheet_size <= 2000:
        return "L"
    return "XL"


def diversity_key(record: dict) -> tuple:
    return (
        record["location"],
        frozenset(record["formula_ref_types"]),
        size_bucket(record["sheet_size"]),
    )


def select_diverse(
    records: list[dict],
    global_location_counts: Counter,
    global_reftype_counts: Counter,
) -> list[dict]:
    if len(records) <= MAX_PER_SKETCH:
        return records

    records_by_diversity = defaultdict(list)
    for record in records:
        records_by_diversity[diversity_key(record)].append(record)

    representatives = []
    for key, group_records in records_by_diversity.items():
        representative = min(
            group_records,
            key=lambda record: abs(record["sheet_size"] - GLOBAL_MEDIAN_SHEET_SIZE),
        )
        representatives.append((key, representative))

    if len(representatives) <= MAX_PER_SKETCH:
        return [record for _, record in representatives]

    def rarity_score(key: tuple) -> int:
        location, ref_types, _ = key
        return global_location_counts.get(location, 0) + global_reftype_counts.get(ref_types, 0)

    representatives.sort(key=lambda item: rarity_score(item[0]))
    return [record for _, record in representatives[:MAX_PER_SKETCH]]


def select_records_by_diversity(records: list[dict]) -> list[dict]:
    records_by_sketch = defaultdict(list)
    global_location_counts = Counter()
    global_reftype_counts = Counter()

    for record in records:
        records_by_sketch[record["sketch"]].append(record)
        global_location_counts[record["location"]] += 1
        global_reftype_counts[frozenset(record["formula_ref_types"])] += 1

    selected_records = []
    for sketch in sorted(records_by_sketch):
        selected_records.extend(
            select_diverse(
                records_by_sketch[sketch],
                global_location_counts,
                global_reftype_counts,
            )
        )
    return selected_records
