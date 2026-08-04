from __future__ import annotations

import json
import re
import threading
import time
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

from fxbench.llm_utils import LLMClient


ERROR_PATTERNS = re.compile(r"#REF!|#NAME\?|#VALUE!|#NULL!|#N/A|#DIV/0!|#NUM!", re.IGNORECASE)

PROMPT_TEMPLATE = """\
You are an expert spreadsheet analyst. You will be given a spreadsheet region \
as a markdown table, a target cell location, and the formula in that cell. \
Analyze the surrounding context and the formula to provide structured annotations.

The target cell is marked with [TARGET] in the table. Formulas in other cells \
are shown as their formula strings (e.g., =B2+C2), not their computed values.

## Spreadsheet Context

{markdown_table}

**Target cell:** {cell_location}
**Formula:** {formula}

## Annotations Required

### 1. Location Classification
Where is the formula situated relative to data structures in the sheet?
- `table_cell` - Inside a structured table or rectangular data region with clear headers
- `table_adjacent` - Within 5 cells of a structured table (e.g., summary row below, total column beside)
- `unstructured` - No obvious tabular structure nearby

### 2. Intent Hint
What contextual clue best reveals the purpose of this formula? Pick the single strongest hint source, unless multiple sources contribute equally.
- `col_header` - A column header above the formula or its referenced cells describes the intent
- `row_label` - A row label to the left of the formula describes the intent
- `adjacent_formula` - Neighboring cells contain similar formulas suggesting a pattern
- `adjacent_values` - A pattern in adjacent cell values hints at the computation
- `combination` - Multiple hint sources contribute equally (specify which ones)
- `no_hint` - No clear contextual clue available to infer what the formula computes

Identify the specific text or pattern that provides the hint.

### 3. Intent Description
In one short sentence, describe what this formula computes based on the surrounding context. \
If context is insufficient, state "unclear from context."

Respond ONLY with the following JSON, no other text:
{{
  "location": "table_cell | table_adjacent | unstructured",
  "intent_hint": {{
    "source": "col_header | row_label | adjacent_formula | adjacent_values | combination | no_hint",
    "hint_text": "the specific text or pattern that provides the hint, or null if no_hint"
  }},
  "intent_description": "one sentence describing what the formula computes"
}}"""


class RateLimiter:
    def __init__(self, rpm: int):
        self.interval = 60.0 / rpm
        self.lock = threading.Lock()
        self.next_allowed = 0.0

    def wait(self) -> None:
        with self.lock:
            now = time.monotonic()
            if now < self.next_allowed:
                delay = self.next_allowed - now
                self.next_allowed += self.interval
            else:
                delay = 0.0
                self.next_allowed = now + self.interval
        if delay > 0:
            time.sleep(delay)


def column_letter(column_index: int) -> str:
    result = ""
    while column_index > 0:
        column_index, remainder = divmod(column_index - 1, 26)
        result = chr(remainder + ord("A")) + result
    return result


def cell_display(cell) -> str:
    if cell.data_type == "f":
        value = cell.value
        if value is None:
            return ""
        if hasattr(value, "text"):
            return f"={value.text}"
        text = str(value)
        if not text.startswith("="):
            return f"={text}"
        return text
    if cell.value is None:
        return ""
    return str(cell.value)


def parse_cell_address(address: str) -> tuple[int, int] | None:
    match = re.match(r"([A-Z]+)(\d+)", address.upper())
    if not match:
        return None
    column_text = match.group(1)
    row = int(match.group(2))
    column = 0
    for char in column_text:
        column = column * 26 + (ord(char) - ord("A") + 1)
    return column, row


def build_markdown_table_for_cell(
    worksheet,
    cell_address: str,
    row_margin: int = 15,
    column_margin: int = 10,
) -> str | None:
    parsed_address = parse_cell_address(cell_address)
    if parsed_address is None:
        return None

    target_column, target_row = parsed_address
    min_row = max(1, target_row - row_margin)
    max_row = min(worksheet.max_row or target_row, target_row + row_margin)
    min_column = max(1, target_column - column_margin)
    max_column = min(worksheet.max_column or target_column, target_column + column_margin)

    column_headers = [column_letter(column) for column in range(min_column, max_column + 1)]
    rows = [
        "| Row | " + " | ".join(column_headers) + " |",
        "|-----|" + "|".join(["---"] * len(column_headers)) + "|",
    ]

    for row_index in range(min_row, max_row + 1):
        values = []
        for column_index in range(min_column, max_column + 1):
            if row_index == target_row and column_index == target_column:
                values.append("[TARGET]")
            else:
                values.append(cell_display(worksheet.cell(row=row_index, column=column_index)))
        rows.append(f"| {row_index} | " + " | ".join(values) + " |")

    return "\n".join(rows)


def missing_markdown_rows(sheet_items: list[tuple[str, list[dict]]]) -> list[dict]:
    rows = []
    for _, cells in sheet_items:
        for cell in cells:
            rows.append({"idx": cell["idx"], "md_table": None})
    return rows


def extract_markdown_tables_from_workbook(
    workbook_path: str,
    sheet_items: list[tuple[str, list[dict]]],
    row_margin: int = 15,
    column_margin: int = 10,
) -> list[dict]:
    if not Path(workbook_path).exists():
        return missing_markdown_rows(sheet_items)

    try:
        workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=False)
    except (BadZipFile, InvalidFileException, KeyError, OSError, ParseError, ValueError):
        return missing_markdown_rows(sheet_items)

    try:
        results = []
        for sheet_name, cells in sheet_items:
            if sheet_name not in workbook.sheetnames:
                results.extend({"idx": cell["idx"], "md_table": None} for cell in cells)
                continue

            worksheet = workbook[sheet_name]
            for cell in cells:
                markdown_table = build_markdown_table_for_cell(
                    worksheet,
                    cell["cell"],
                    row_margin=row_margin,
                    column_margin=column_margin,
                )
                results.append({"idx": cell["idx"], "md_table": markdown_table})
        return results
    finally:
        workbook.close()


def normalized_formula(formula: str) -> str:
    if formula.startswith("="):
        return formula
    return f"={formula}"


def build_annotation_prompt(markdown_table: str, cell: str, formula: str) -> str:
    return PROMPT_TEMPLATE.format(
        markdown_table=markdown_table,
        cell_location=cell,
        formula=normalized_formula(formula),
    )


def skip_reason_for_error(error: Exception) -> str:
    error_text = str(error).lower()
    if any(text in error_text for text in ("429", "rate limit", "ratelimit", "too many requests")):
        return "rate_limit"
    if any(
        text in error_text
        for text in (
            "model_error",
            "invalid content",
            "content_filter",
            "content management policy",
            "harmful content",
            "responsibleaipolicyviolation",
            "filtered",
        )
    ):
        return "content_filter"
    if any(text in error_text for text in ("credential", "get_token", "azure cli", "authentication")):
        return "auth_error"
    return f"llm_error:{type(error).__name__}"


def call_llm_for_annotation(
    client: LLMClient,
    model: str,
    markdown_table: str,
    cell: str,
    formula: str,
) -> dict:
    prompt = build_annotation_prompt(markdown_table, cell, formula)
    try:
        result = client.prompt_llm(model=model, question=prompt, max_tokens=500, retries=0, retry_delay=3)
    except Exception as error:
        return {"skipped": True, "reason": skip_reason_for_error(error)}

    raw_response = result.get("response", "")
    clean_response = client.extract_json_from_response(raw_response)
    try:
        return json.loads(clean_response)
    except json.JSONDecodeError:
        return {"parse_error": True, "reason": "parse_error"}
