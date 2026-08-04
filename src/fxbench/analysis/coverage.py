"""Required-vs-exposed theoretical limit analysis.

This analysis separates formula-reference cells from intent cells and measures
how much of each subset is printed or mentioned by each context strategy.

Usage:
    python -m analysis.required_exposed_theoretical_limit [--take N]
"""

from __future__ import annotations

import argparse
import csv
import json
from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

from fxbench.analysis import context_evidence as base
from fxbench.evaluation.formula_sketch import formula_to_metric_sketch


@dataclass(frozen=True)
class StrategySpec:
    name: str
    label: str
    evidence_fn: Callable


@dataclass
class ProblemAccumulator:
    covered_problems: int = 0
    total_problems: int = 0

    def add(self, covered: bool) -> None:
        self.covered_problems += int(covered)
        self.total_problems += 1

    def payload(self) -> dict:
        pct = (
            self.covered_problems / self.total_problems * 100
            if self.total_problems
            else None
        )
        return {
            "covered_problems": self.covered_problems,
            "total_problems": self.total_problems,
            "pct": pct,
        }


METRIC_NAMES = (
    "intent_printed",
    "intent_mentioned",
    "reference_printed",
    "reference_mentioned",
    "sibling_formula_printed",
    "theoretical_upper_limit",
    "theoretical_upper_limit_v3",
    "theoretical_upper_limit_v4",
)


def _filter_required_cells(cells: set[str], ws, target_cell: str) -> set[str]:
    required: set[str] = set()
    for addr in cells:
        if addr == target_cell:
            continue
        r, c = coordinate_to_tuple(addr)
        cell = ws.cell(row=r, column=c)
        if cell.value not in {None, ""}:
            required.add(addr)
    return required


def _intent_cells_from_fields(intent_data: dict | None, fields: tuple[str, ...]) -> set[str]:
    cells: set[str] = set()
    if not intent_data:
        return cells
    for field in fields:
        for ref in intent_data.get(field, []) or []:
            if not isinstance(ref, str):
                continue
            ref = base._normalize_ref(ref.strip().upper())
            if not ref:
                continue
            try:
                cells.update(base._expand_ref(ref))
            except ValueError:
                continue
    return cells


def _formula_cell_value(cell) -> str | None:
    raw = cell.value
    data_type = getattr(cell, "data_type", "n")
    if data_type == "f":
        if isinstance(raw, ArrayFormula):
            formula = raw.text
        elif isinstance(raw, DataTableFormula):
            return None
        elif raw is None:
            return None
        else:
            formula = str(raw)
    elif isinstance(raw, str) and raw.startswith("="):
        formula = raw
    else:
        return None
    if not formula:
        return None
    return formula if formula.startswith("=") else f"={formula}"


def _sibling_formula_cells(
    ws,
    target_cell: str,
    target_formula: str,
) -> tuple[set[str], str]:
    target_sketch = formula_to_metric_sketch(target_formula)
    if not target_sketch:
        return set(), target_sketch

    siblings: set[str] = set()
    for row in ws.iter_rows(
        min_row=ws.min_row,
        max_row=ws.max_row,
        min_col=ws.min_column,
        max_col=ws.max_column,
    ):
        for cell in row:
            addr = f"{get_column_letter(cell.column)}{cell.row}"
            if addr == target_cell:
                continue
            formula = _formula_cell_value(cell)
            if formula is None:
                continue
            if formula_to_metric_sketch(formula) == target_sketch:
                siblings.add(addr)
    return siblings, target_sketch


def _windowed_evidence(row_margin: int, col_margin: int) -> Callable:
    def evidence(ws, target_cell: str, wb_path: str | None = None):
        del wb_path
        visible = base._visible_windowed(
            ws,
            target_cell,
            row_margin=row_margin,
            col_margin=col_margin,
        )
        return base._all_printed_evidence_with_inferred_types(ws, target_cell, visible)

    return evidence


def _strategy_specs(
    minisheet_width: int = base.DEFAULT_MINISHEET_WIDTH,
    minisheet_kwargs: dict | None = None,
) -> dict[str, StrategySpec]:
    evidence_fns = base._evidence_fns(
        minisheet_width=minisheet_width,
        minisheet_kwargs=minisheet_kwargs,
    )
    specs: dict[str, StrategySpec] = {}
    for strategy in base.STRATEGIES:
        specs[strategy] = StrategySpec(
            name=strategy,
            label=base.STRATEGY_LABELS[strategy],
            evidence_fn=evidence_fns[strategy],
        )

    specs["Windowed@30r20c"] = StrategySpec(
        name="Windowed@30r20c",
        label="Windowed (±30r/±20c)",
        evidence_fn=_windowed_evidence(row_margin=30, col_margin=20),
    )
    specs["Windowed@40r30c"] = StrategySpec(
        name="Windowed@40r30c",
        label="Windowed (±40r/±30c)",
        evidence_fn=_windowed_evidence(row_margin=40, col_margin=30),
    )
    return specs


def _default_strategy_order() -> list[str]:
    order: list[str] = []
    for strategy in base.STRATEGIES:
        order.append(strategy)
        if strategy == "Windowed":
            order.extend(["Windowed@30r20c", "Windowed@40r30c"])
    return order


def _new_accumulators(strategies: list[str]) -> dict[str, dict[str, ProblemAccumulator]]:
    return {
        strategy: {metric: ProblemAccumulator() for metric in METRIC_NAMES}
        for strategy in strategies
    }


def _is_abstain_record(record: dict) -> bool:
    return str(record.get("function", "")).upper() == "ABSTAIN"


def compute_required_exposed_limits(
    records: list[dict],
    *,
    strategies: list[str] | None = None,
    intent_cells_by_id: dict[str, dict] | None = None,
    minisheet_width: int = base.DEFAULT_MINISHEET_WIDTH,
    minisheet_kwargs: dict | None = None,
    include_details: bool = False,
) -> tuple[dict[str, dict], dict]:
    specs = _strategy_specs(
        minisheet_width=minisheet_width,
        minisheet_kwargs=minisheet_kwargs,
    )
    selected_strategies = strategies or _default_strategy_order()
    missing = [strategy for strategy in selected_strategies if strategy not in specs]
    if missing:
        raise ValueError(f"Unknown strategies: {missing}")

    intent_cells_by_id = (
        base.load_intent_annotations()
        if intent_cells_by_id is None
        else intent_cells_by_id
    )
    accumulators = _new_accumulators(selected_strategies)
    total = 0
    errors = 0
    cross_sheet_count = 0
    outside_reference_cells = 0
    outside_intent_cells = 0
    reference_cell_total = 0
    intent_cell_total = 0
    sibling_formula_cell_total = 0
    required_cell_total = 0
    abstain_count = 0
    details: list[dict] = []

    for index, rec in enumerate(records):
        if _is_abstain_record(rec):
            row = {
                "id": rec["id"],
                "abstain": True,
                "reference_required": 0,
                "intent_required": 0,
                "direct_intent_required": 0,
                "data_description_required": 0,
                "sibling_formula_required": 0,
                "required": 0,
                "cross_sheet": False,
                "outside_reference_cells": 0,
                "outside_intent_cells": 0,
            }
            for strategy in selected_strategies:
                for metric in METRIC_NAMES:
                    accumulators[strategy][metric].add(True)
                if include_details:
                    row[strategy] = {
                        metric: {
                            "covered": True,
                            "abstain_default": True,
                            "missing": [],
                        }
                        for metric in METRIC_NAMES
                    }
            details.append(row)
            abstain_count += 1
            total += 1
            continue

        formula = rec["formula"]
        target_cell = rec["cell"]
        wb_path = Path(rec["wb_local_path"])
        if not wb_path.is_absolute():
            wb_path = base.REPO_ROOT / wb_path

        referenced, has_cross = base.extract_required_refs(formula)
        if has_cross:
            cross_sheet_count += 1
        intent_data = intent_cells_by_id.get(rec["id"])
        direct_intent = _intent_cells_from_fields(intent_data, ("direct_intent",))
        data_description = _intent_cells_from_fields(intent_data, ("data_description",))
        intent = direct_intent | data_description

        try:
            wb = load_workbook(str(wb_path), read_only=False, data_only=False)
            ws = wb[rec["sheet"]]
        except Exception as exc:
            print(f"  [{index + 1}] ERROR: {exc}")
            errors += 1
            continue

        try:
            base._MINISHEET_CELL_CACHE.clear()
            used_range = base._visible_full_sheet(ws)

            reference_raw = referenced & used_range
            intent_raw = intent & used_range
            direct_intent_raw = direct_intent & used_range
            data_description_raw = data_description & used_range
            sibling_formula_raw, sibling_formula_sketch = _sibling_formula_cells(
                ws,
                target_cell,
                formula,
            )
            sibling_formula_raw &= used_range
            outside_reference_cells += len(referenced - used_range)
            outside_intent_cells += len(intent - used_range)

            reference_required = _filter_required_cells(reference_raw, ws, target_cell)
            sibling_formula_required = _filter_required_cells(
                sibling_formula_raw,
                ws,
                target_cell,
            )
            direct_intent_required = _filter_required_cells(
                direct_intent_raw,
                ws,
                target_cell,
            )
            data_description_required = _filter_required_cells(
                data_description_raw,
                ws,
                target_cell,
            )
            intent_required = _filter_required_cells(intent_raw, ws, target_cell)
            required = reference_required | intent_required

            reference_cell_total += len(reference_required)
            intent_cell_total += len(intent_required)
            sibling_formula_cell_total += len(sibling_formula_required)
            required_cell_total += len(required)

            row = {
                "id": rec["id"],
                "reference_required": len(reference_required),
                "intent_required": len(intent_required),
                "direct_intent_required": len(direct_intent_required),
                "data_description_required": len(data_description_required),
                "sibling_formula_required": len(sibling_formula_required),
                "required": len(required),
                "cross_sheet": has_cross,
                "outside_reference_cells": len(referenced - used_range),
                "outside_intent_cells": len(intent - used_range),
            }

            for strategy in selected_strategies:
                evidence = specs[strategy].evidence_fn(ws, target_cell, str(wb_path))
                metric_sets = {
                    "intent_printed": (evidence.printed, intent_required),
                    "intent_mentioned": (evidence.mentioned, intent_required),
                    "reference_printed": (evidence.printed, reference_required),
                    "reference_mentioned": (evidence.mentioned, reference_required),
                    "sibling_formula_printed": (
                        evidence.printed,
                        sibling_formula_required,
                    ),
                }
                for metric, (covered, required_set) in metric_sets.items():
                    accumulators[strategy][metric].add(required_set.issubset(covered))
                upper_mentioned_required = data_description_required | reference_required
                upper_success = (
                    direct_intent_required.issubset(evidence.printed)
                    and upper_mentioned_required.issubset(evidence.mentioned)
                )
                upper_v3_printed_required = direct_intent_required | sibling_formula_required
                upper_v3_success = (
                    upper_v3_printed_required.issubset(evidence.printed)
                    and upper_mentioned_required.issubset(evidence.mentioned)
                )
                upper_v4_mentioned_required = (
                    upper_mentioned_required | sibling_formula_required
                )
                upper_v4_success = (
                    direct_intent_required.issubset(evidence.printed)
                    and upper_v4_mentioned_required.issubset(evidence.mentioned)
                )
                accumulators[strategy]["theoretical_upper_limit"].add(upper_success)
                accumulators[strategy]["theoretical_upper_limit_v3"].add(upper_v3_success)
                accumulators[strategy]["theoretical_upper_limit_v4"].add(upper_v4_success)

                if include_details:
                    row[strategy] = {
                        metric: {
                            "covered": required_set.issubset(covered),
                            "covered_cells": len(covered & required_set),
                            "required_cells": len(required_set),
                            "missing": sorted(required_set - covered),
                        }
                        for metric, (covered, required_set) in metric_sets.items()
                    }
                    row[strategy]["theoretical_upper_limit"] = {
                        "covered": upper_success,
                        "total": 1,
                        "missing": sorted(
                            (direct_intent_required - evidence.printed)
                            | (upper_mentioned_required - evidence.mentioned)
                        ),
                        "direct_intent_printed_missing": sorted(
                            direct_intent_required - evidence.printed
                        ),
                        "mentioned_missing": sorted(
                            upper_mentioned_required - evidence.mentioned
                        ),
                    }
                    row[strategy]["theoretical_upper_limit_v3"] = {
                        "covered": upper_v3_success,
                        "total": 1,
                        "missing": sorted(
                            (upper_v3_printed_required - evidence.printed)
                            | (upper_mentioned_required - evidence.mentioned)
                        ),
                        "direct_intent_or_sibling_printed_missing": sorted(
                            upper_v3_printed_required - evidence.printed
                        ),
                        "mentioned_missing": sorted(
                            upper_mentioned_required - evidence.mentioned
                        ),
                    }
                    row[strategy]["theoretical_upper_limit_v4"] = {
                        "covered": upper_v4_success,
                        "total": 1,
                        "missing": sorted(
                            (direct_intent_required - evidence.printed)
                            | (upper_v4_mentioned_required - evidence.mentioned)
                        ),
                        "direct_intent_printed_missing": sorted(
                            direct_intent_required - evidence.printed
                        ),
                        "mentioned_missing": sorted(
                            upper_v4_mentioned_required - evidence.mentioned
                        ),
                    }
                    row["sibling_formula_cells"] = sorted(sibling_formula_required)
                    row["sibling_formula_sketch"] = sibling_formula_sketch

            details.append(row)
            total += 1
        finally:
            base._MINISHEET_CELL_CACHE.clear()
            wb.close()

    payload: dict[str, dict] = {}
    for strategy in selected_strategies:
        strategy_payload = {
            metric: accumulators[strategy][metric].payload()
            for metric in METRIC_NAMES
        }
        strategy_payload.update(
            {
                "label": specs[strategy].label,
                "problems": total,
                "reference_cells": reference_cell_total,
                "intent_cells": intent_cell_total,
                "sibling_formula_cells": sibling_formula_cell_total,
                "required_cells": required_cell_total,
            }
        )
        payload[strategy] = strategy_payload

    meta = {
        "total": total,
        "errors": errors,
        "cross_sheet_count": cross_sheet_count,
        "outside_reference_cells": outside_reference_cells,
        "outside_intent_cells": outside_intent_cells,
        "reference_cells": reference_cell_total,
        "intent_cells": intent_cell_total,
        "sibling_formula_cells": sibling_formula_cell_total,
        "required_cells": required_cell_total,
        "abstain_count": abstain_count,
        "details": details if include_details else [],
    }
    return payload, meta


def _fmt_pct(value: float | None) -> str:
    return "n/a" if value is None else f"{value:5.1f}%"


def _write_csv(path: Path, payload: dict[str, dict]) -> None:
    fields = [
        "strategy",
        "label",
        "intent_printed_records",
        "intent_mentioned_records",
        "reference_printed_records",
        "reference_mentioned_records",
        "sibling_formula_printed_records",
        "theoretical_upper_limit_records",
        "theoretical_upper_limit_v3_records",
        "theoretical_upper_limit_v4_records",
        "total_records",
        "intent_printed_pct",
        "intent_mentioned_pct",
        "reference_printed_pct",
        "reference_mentioned_pct",
        "sibling_formula_printed_pct",
        "theoretical_upper_limit_pct",
        "theoretical_upper_limit_v3_pct",
        "theoretical_upper_limit_v4_pct",
        "intent_cells",
        "reference_cells",
        "sibling_formula_cells",
        "required_cells",
    ]
    with path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=fields)
        writer.writeheader()
        for strategy, metrics in payload.items():
            writer.writerow(
                {
                    "strategy": strategy,
                    "label": metrics["label"],
                    "intent_printed_records": metrics["intent_printed"]["covered_problems"],
                    "intent_mentioned_records": metrics["intent_mentioned"]["covered_problems"],
                    "reference_printed_records": metrics["reference_printed"]["covered_problems"],
                    "reference_mentioned_records": metrics["reference_mentioned"]["covered_problems"],
                    "sibling_formula_printed_records": metrics["sibling_formula_printed"]["covered_problems"],
                    "theoretical_upper_limit_records": metrics["theoretical_upper_limit"]["covered_problems"],
                    "theoretical_upper_limit_v3_records": metrics["theoretical_upper_limit_v3"]["covered_problems"],
                    "theoretical_upper_limit_v4_records": metrics["theoretical_upper_limit_v4"]["covered_problems"],
                    "total_records": metrics["problems"],
                    "intent_printed_pct": metrics["intent_printed"]["pct"],
                    "intent_mentioned_pct": metrics["intent_mentioned"]["pct"],
                    "reference_printed_pct": metrics["reference_printed"]["pct"],
                    "reference_mentioned_pct": metrics["reference_mentioned"]["pct"],
                    "sibling_formula_printed_pct": metrics["sibling_formula_printed"]["pct"],
                    "theoretical_upper_limit_pct": metrics["theoretical_upper_limit"]["pct"],
                    "theoretical_upper_limit_v3_pct": metrics["theoretical_upper_limit_v3"]["pct"],
                    "theoretical_upper_limit_v4_pct": metrics["theoretical_upper_limit_v4"]["pct"],
                    "intent_cells": metrics["intent_cells"],
                    "reference_cells": metrics["reference_cells"],
                    "sibling_formula_cells": metrics["sibling_formula_cells"],
                    "required_cells": metrics["required_cells"],
                }
            )


def print_summary(payload: dict[str, dict], meta: dict) -> None:
    print()
    print(
        f"Processed {meta['total']} problems "
        f"({meta['errors']} errors, {meta['cross_sheet_count']} with cross-sheet refs)"
    )
    print()
    header = (
        f"{'Strategy':<34}"
        f"{'IntentPr':>10}"
        f"{'IntentMen':>10}"
        f"{'RefPr':>10}"
        f"{'RefMen':>10}"
        f"{'SibPr':>10}"
        f"{'UpperV2':>10}"
        f"{'UpperV3':>10}"
        f"{'UpperV4':>10}"
    )
    print(header)
    print("-" * len(header))
    for metrics in payload.values():
        print(
            f"{metrics['label']:<34}"
            f"{metrics['intent_printed']['covered_problems']:>10}"
            f"{metrics['intent_mentioned']['covered_problems']:>10}"
            f"{metrics['reference_printed']['covered_problems']:>10}"
            f"{metrics['reference_mentioned']['covered_problems']:>10}"
            f"{metrics['sibling_formula_printed']['covered_problems']:>10}"
            f"{metrics['theoretical_upper_limit']['covered_problems']:>10}"
            f"{metrics['theoretical_upper_limit_v3']['covered_problems']:>10}"
            f"{metrics['theoretical_upper_limit_v4']['covered_problems']:>10}"
        )
