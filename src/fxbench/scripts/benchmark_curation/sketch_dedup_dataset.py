import re
from pathlib import Path
from collections import defaultdict
from datasets import load_dataset
from rich.progress import (
    BarColumn, MofNCompleteColumn, Progress, SpinnerColumn,
    TextColumn, TimeElapsedColumn, TimeRemainingColumn
)
import json
from io import BytesIO
import openpyxl
from fxbench.utils import read_jsonl


INPUT_DIR = Path("Assets/records/valid_records")
HF_DATASET = "tianzl66/Sheetpedia_xlsx"
WORKBOOK_DIR = Path("Assets/workbooks")

def formula_to_sketch(formula: str) -> str:
    """Convert formula to sketch by replacing references, numbers, strings."""
    s = formula
    if s.startswith("{=") and s.endswith("}"):
        s = s[1:-1]
    s = re.sub(r'"[^"]*"', "<string>", s)
    s = re.sub(r"'[^']+'\!", "SHEET!", s)
    s = re.sub(r"[A-Za-z_]\w*\!", "SHEET!", s)
    s = re.sub(r"[A-Za-z_]\w*\[@\[[^\]]*\]\]", "TABLE[@COL]", s)
    s = re.sub(r"[A-Za-z_]\w*\[\[[^\]]*\]\]", "TABLE[COL]", s)
    s = re.sub(r"[A-Za-z_]\w*\[@[^\]]*\]", "TABLE[@COL]", s)
    s = re.sub(r"[A-Za-z_]\w*\[[^\]]*\]", "TABLE[COL]", s)
    s = re.sub(r"\$?[A-Z]{1,3}\$?\d{1,7}:\$?[A-Z]{1,3}\$?\d{1,7}", "<range>", s)
    s = re.sub(r"\$?[A-Z]{1,3}\$?\d{1,7}", "<cell>", s)
    s = re.sub(r"(?<![A-Za-z])(TRUE|FALSE)(?![A-Za-z])", "<bool>", s, flags=re.IGNORECASE)
    s = re.sub(r"\b\d+\.\d+\b", "<float>", s)
    s = re.sub(r"\b\d+\b", "<int>", s)
    return s


def get_formula_text(cell) -> str | None:
    if cell.data_type != "f":
        return None
    raw = cell.value
    if hasattr(raw, "text"):
        text = raw.text
    elif isinstance(raw, str):
        text = raw
    else:
        return None
    if not text:
        return None
    if not text.startswith("="):
        text = "=" + text
    return text

def extract_primary_function(formula: str) -> str:
    s = formula.lstrip("=")
    m = re.match(r"[\s]*([A-Za-z][A-Za-z0-9_.]*?)\s*\(", s)
    if m:
        return m.group(1).upper()
    return "EXPRESSION"


def main():
    valid_files = read_jsonl(INPUT_DIR / "function_level.jsonl")
    ds = load_dataset(HF_DATASET)["train"]
    
    sketch_map: dict[str, dict[str, list[dict]]] = defaultdict(lambda: defaultdict(list))
    seen: set[tuple[str, int]] = set()
    total_formulas_before = 0
    total_formulas_after = 0
    
    with Progress(
        SpinnerColumn(), TextColumn("[bold blue]{task.description}"), 
        BarColumn(), MofNCompleteColumn(), TimeElapsedColumn(), TimeRemainingColumn()
    ) as progress:
        task = progress.add_task("Extracting sketches", total=len(valid_files))
        
        for file_info in valid_files:
            index = file_info["index"]
            key = file_info["key"]
            xlsx_bytes = ds[index]["xlsx"] 
            
            try:
                wb = openpyxl.load_workbook(BytesIO(xlsx_bytes), read_only=True, data_only=False)
                workbook_saved = False 
                
                for sheet_info in file_info["valid_sheets"]:
                    sheet_name = sheet_info["sheet_name"] 
                    
                    if sheet_name not in wb.sheetnames:
                        continue 
                    
                    ws = wb[sheet_name]
                    for row in ws.iter_rows():
                        for cell in row:
                            formula_text = get_formula_text(cell)
                            if formula_text is None:
                                continue
                            
                            total_formulas_before += 1 
                            sketch = formula_to_sketch(formula_text)
                            func = extract_primary_function(formula_text)
                            dedup_key = (sketch, index)
                            
                            if dedup_key in seen:
                                continue
                            
                            seen.add(dedup_key)
                            total_formulas_after += 1
                            wb_local_path = f"{WORKBOOK_DIR}/{index}.xlsx"
                            
                            record = {
                                "hf_index": index,
                                "sheet": sheet_name,
                                "cell": cell.coordinate,
                                "formula": formula_text,
                                "id": f"{key}__{sheet_name}__{cell.coordinate}",
                                "wb_local_path": wb_local_path
                            }
                            sketch_map[func][sketch].append(record)
                            
                            if not workbook_saved:
                                if not WORKBOOK_DIR.exists():
                                    WORKBOOK_DIR.mkdir(parents=True)
                                wb.save(wb_local_path)
                                workbook_saved = True
                
                wb.close()
            
            except Exception as e:
                print(f"Error processing {key}: {e}")
                continue
            
            progress.advance(task)
    
    out_path = INPUT_DIR / "sketch_map_deduped.json" 
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump({func: dict(sketches) for func, sketches in sketch_map.items()}, f, indent=2)
    
    total_sketches = sum(len(s) for s in sketch_map.values())
    print("Sketch extraction completed.")
    print(f"Total formulas before deduplication: {total_formulas_before}")
    print(f"Total formulas after deduplication: {total_formulas_after}")
    print(f"Total sketches extracted: {total_sketches}")
    print(f"Primary functions found: {len(sketch_map)}")
    
if __name__ == "__main__":
    main()
    