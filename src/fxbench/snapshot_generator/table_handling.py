from __future__ import annotations

from pathlib import Path

import numpy as np
import openpyxl
from openpyxl import load_workbook


def get_excel_range(min_row: int, min_col: int, max_row: int, max_col: int) -> str:
    start_column = openpyxl.utils.get_column_letter(min_col)
    end_column = openpyxl.utils.get_column_letter(max_col)
    return f"{start_column}{min_row}:{end_column}{max_row}"


def parse_cell_reference(cell_ref: str) -> tuple[int, int]:
    row_index, column_index = openpyxl.utils.coordinate_to_tuple(cell_ref)
    return row_index, column_index


def extract_excel_table(sheet) -> list[dict]:
    tables = []
    for table in sheet.tables.values():
        start_cell, end_cell = table.ref.split(":")
        min_row, min_col = parse_cell_reference(start_cell)
        max_row, max_col = parse_cell_reference(end_cell)

        tables.append(
            {
                "min_row": min_row,
                "max_row": max_row,
                "min_col": min_col,
                "max_col": max_col,
                "type": "excel_table",
                "name": table.name,
                "display_name": table.displayName,
            }
        )
    return tables


def create_occupancy_matrix(sheet, max_row: int, max_col: int) -> np.ndarray:
    occupancy = np.zeros((max_row, max_col), dtype=int)
    for row_index in range(1, max_row + 1):
        for column_index in range(1, max_col + 1):
            if sheet.cell(row_index, column_index).value is not None:
                occupancy[row_index - 1, column_index - 1] = 1
    return occupancy


def mask_detected_regions(occupancy: np.ndarray, regions: list[dict]) -> np.ndarray:
    masked = occupancy.copy()
    for region in regions:
        row_start = region["min_row"] - 1
        row_end = region["max_row"] - 1
        column_start = region["min_col"] - 1
        column_end = region["max_col"] - 1
        masked[row_start : row_end + 1, column_start : column_end + 1] = 0
    return masked


def find_dense_ranges(
    fill_rates: np.ndarray,
    threshold: float,
    max_empty_gap: int = 2,
) -> list[tuple[int, int]]:
    ranges = []
    start = None
    empty_gap_count = 0

    for index, rate in enumerate(fill_rates):
        if rate >= threshold:
            if start is None:
                start = index
            empty_gap_count = 0
        elif rate > 0:
            if start is None:
                start = index
            empty_gap_count = 0
        else:
            if start is not None:
                empty_gap_count += 1
                if empty_gap_count > max_empty_gap:
                    ranges.append((start, index - empty_gap_count))
                    start = None
                    empty_gap_count = 0

    if start is not None:
        end = len(fill_rates) - 1 - empty_gap_count if empty_gap_count > 0 else len(fill_rates) - 1
        ranges.append((start, end))

    return ranges


def merge_overlapping_regions(regions: list[dict]) -> list[dict]:
    if not regions:
        return []

    regions = sorted(regions, key=lambda region: (region["min_row"], region["min_col"]))
    merged = [regions[0]]

    for current in regions[1:]:
        last = merged[-1]
        row_overlap = current["min_row"] <= last["max_row"] and current["max_row"] >= last["min_row"]
        column_overlap = current["min_col"] <= last["max_col"] and current["max_col"] >= last["min_col"]

        if row_overlap and column_overlap:
            merged[-1] = {
                "min_row": min(last["min_row"], current["min_row"]),
                "max_row": max(last["max_row"], current["max_row"]),
                "min_col": min(last["min_col"], current["min_col"]),
                "max_col": max(last["max_col"], current["max_col"]),
                "type": last["type"],
            }
        else:
            merged.append(current)

    return merged


def detect_heuristic_regions(
    sheet,
    occupancy: np.ndarray,
    fill_threshold: float,
    min_region_size: int,
    max_gap: int,
) -> list[dict]:
    rows, columns = occupancy.shape
    row_fill_rates = occupancy.sum(axis=1) / columns
    column_fill_rates = occupancy.sum(axis=0) / rows

    dense_row_ranges = find_dense_ranges(row_fill_rates, fill_threshold, max_gap)
    dense_column_ranges = find_dense_ranges(column_fill_rates, fill_threshold, max_gap)

    regions = []
    for row_start, row_end in dense_row_ranges:
        for column_start, column_end in dense_column_ranges:
            region_data = occupancy[row_start : row_end + 1, column_start : column_end + 1]
            region_fill_rate = region_data.sum() / region_data.size
            if region_fill_rate < fill_threshold:
                continue

            region_size = (row_end - row_start + 1) * (column_end - column_start + 1)
            if region_size < min_region_size:
                continue

            regions.append(
                {
                    "min_row": row_start + 1,
                    "max_row": row_end + 1,
                    "min_col": column_start + 1,
                    "max_col": column_end + 1,
                    "type": "heuristic_table",
                    "fill_rate": float(region_fill_rate),
                    "size": int(region_size),
                }
            )

    return merge_overlapping_regions(regions)


def detect_regions_from_excel(filepath: str | Path, sheet_name: str | None = None) -> list[dict]:
    workbook = load_workbook(filename=filepath, data_only=True)
    try:
        sheet = workbook[sheet_name] if sheet_name else workbook.active
        max_row = sheet.max_row
        max_col = sheet.max_column

        regions = []
        excel_tables = extract_excel_table(sheet)
        regions.extend(excel_tables)

        occupancy_matrix = create_occupancy_matrix(sheet, max_row, max_col)
        masked_occupancy = mask_detected_regions(occupancy_matrix, excel_tables)
        heuristic_regions = detect_heuristic_regions(
            sheet,
            masked_occupancy,
            fill_threshold=0.3,
            min_region_size=2,
            max_gap=2,
        )
        regions.extend(heuristic_regions)

        regions = sorted(regions, key=lambda region: (region["min_row"], region["min_col"]))
        for index, region in enumerate(regions):
            region["region_id"] = index + 1
            region["excel_range"] = get_excel_range(
                region["min_row"],
                region["min_col"],
                region["max_row"],
                region["max_col"],
            )

        return regions
    finally:
        workbook.close()
