from __future__ import annotations

import datetime
import re
from copy import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import range_boundaries
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string


def ensure_output_dir(output_dir: str | Path) -> Path:
    output_path = Path(output_dir)
    output_path.mkdir(parents=True, exist_ok=True)
    return output_path


def safe_filename_part(value: str) -> str:
    return re.sub(r"[^\w.-]+", "_", value)


def output_workbook_path(source_file_path: str | Path, sheet_name: str, suffix: str, output_dir: str | Path) -> Path:
    output_path = ensure_output_dir(output_dir)
    source_path = Path(source_file_path)
    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"{source_path.stem}_{safe_filename_part(sheet_name)}_{suffix}_{timestamp}.xlsx"
    return output_path / filename


def copy_default_font(source_workbook, target_workbook) -> None:
    source_normal = next((style for style in source_workbook._named_styles if style.name == "Normal"), None)
    target_normal = next((style for style in target_workbook._named_styles if style.name == "Normal"), None)
    if source_normal is not None and target_normal is not None:
        target_normal.font = copy(source_normal.font)


def copy_cell(source_cell, target_cell) -> None:
    target_cell.value = source_cell.value
    if source_cell.has_style:
        target_cell.font = copy(source_cell.font)
        target_cell.border = copy(source_cell.border)
        target_cell.fill = copy(source_cell.fill)
        target_cell.number_format = source_cell.number_format
        target_cell.protection = copy(source_cell.protection)
        target_cell.alignment = copy(source_cell.alignment)


def copy_dimensions(source_worksheet, target_worksheet) -> None:
    for column_name in source_worksheet.column_dimensions:
        target_worksheet.column_dimensions[column_name].width = source_worksheet.column_dimensions[column_name].width
    for row_index in source_worksheet.row_dimensions:
        target_worksheet.row_dimensions[row_index].height = source_worksheet.row_dimensions[row_index].height
    if source_worksheet.freeze_panes:
        target_worksheet.freeze_panes = source_worksheet.freeze_panes


def copy_other_sheets(source_workbook, target_workbook, skip_sheet: str) -> None:
    for sheet_name in source_workbook.sheetnames:
        if sheet_name == skip_sheet:
            continue

        source_worksheet = source_workbook[sheet_name]
        target_worksheet = target_workbook.create_sheet(title=sheet_name)
        for row in source_worksheet.iter_rows():
            for source_cell in row:
                target_cell = target_worksheet.cell(row=source_cell.row, column=source_cell.column)
                copy_cell(source_cell, target_cell)
        copy_dimensions(source_worksheet, target_worksheet)


def copy_region_cells(source_worksheet, target_worksheet, region_range: str) -> None:
    min_column, min_row, max_column, max_row = range_boundaries(region_range)
    if min_row > max_row or min_column > max_column:
        raise ValueError(f"Invalid region range: {region_range}")

    for row_index in range(min_row, max_row + 1):
        for column_index in range(min_column, max_column + 1):
            source_cell = source_worksheet.cell(row=row_index, column=column_index)
            target_cell = target_worksheet.cell(row=row_index, column=column_index)
            copy_cell(source_cell, target_cell)


def write_region_summary(worksheet, focused_regions: dict[int, dict]) -> None:
    worksheet.cell(row=1, column=1, value="Region ID")
    worksheet.cell(row=1, column=2, value="Type")
    worksheet.cell(row=1, column=3, value="Range")
    worksheet.cell(row=1, column=4, value="Dependencies")

    for row_index, (region_id, region_data) in enumerate(focused_regions.items(), start=2):
        worksheet.cell(row=row_index, column=1, value=region_id)
        worksheet.cell(row=row_index, column=2, value=region_data.get("type", "unknown"))
        worksheet.cell(row=row_index, column=3, value=region_data.get("range", ""))
        dependencies = region_data.get("dependencies", [])
        worksheet.cell(row=row_index, column=4, value=str(dependencies) if dependencies else "None")


def set_automatic_calculation(workbook) -> None:
    if hasattr(workbook, "calculation") and workbook.calculation is not None:
        workbook.calculation.calcMode = "auto"


def reconstruct_sheet_with_focused_regions(
    source_file_path: str | Path,
    sheet_name: str,
    focused_regions: dict[int, dict],
    output_dir: str | Path,
    *,
    include_summary_sheet: bool = True,
) -> Path:
    output_path = output_workbook_path(source_file_path, sheet_name, "exact_positions", output_dir)
    source_workbook = load_workbook(source_file_path, data_only=False)
    target_workbook = Workbook()

    try:
        if sheet_name not in source_workbook.sheetnames:
            raise ValueError(f"Sheet '{sheet_name}' not found. Available: {source_workbook.sheetnames}")

        source_worksheet = source_workbook[sheet_name]
        copy_default_font(source_workbook, target_workbook)
        target_worksheet = target_workbook.active
        target_worksheet.title = sheet_name

        for region_id, region_data in focused_regions.items():
            region_range = region_data.get("range")
            if not region_range:
                raise ValueError(f"Region {region_id} has no range")
            copy_region_cells(source_worksheet, target_worksheet, region_range)

        copy_dimensions(source_worksheet, target_worksheet)
        copy_other_sheets(source_workbook, target_workbook, skip_sheet=sheet_name)

        if include_summary_sheet:
            summary_sheet = target_workbook.create_sheet("Region_Summary")
            write_region_summary(summary_sheet, focused_regions)

        set_automatic_calculation(target_workbook)
        target_workbook.save(output_path)
        return output_path
    finally:
        source_workbook.close()
        target_workbook.close()


def cells_for_single_cell(cell_reference: str) -> set[tuple[int, int]]:
    column_text, row_index = coordinate_from_string(cell_reference)
    column_index = column_index_from_string(column_text)
    return {(row_index, column_index)}


def cells_for_range(range_text: str, max_row: int, max_column: int) -> set[tuple[int, int]]:
    if ":" not in range_text:
        return cells_for_single_cell(range_text)

    min_column, min_row, max_column_range, max_row_range = range_boundaries(range_text)
    effective_max_row = min(max_row_range, max_row)
    effective_max_column = min(max_column_range, max_column)

    cells = set()
    for row_index in range(min_row, effective_max_row + 1):
        for column_index in range(min_column, effective_max_column + 1):
            cells.add((row_index, column_index))
    return cells


def cells_for_ranges(ranges_to_delete: list[str], max_row: int, max_column: int) -> set[tuple[int, int]]:
    cells = set()
    for range_text in ranges_to_delete:
        try:
            cells.update(cells_for_range(range_text, max_row, max_column))
        except Exception as error:
            raise ValueError(f"Could not parse range '{range_text}': {error}") from error
    return cells


def write_deletion_summary(worksheet, ranges_to_delete: list[str], deleted_cell_count: int) -> None:
    worksheet.cell(row=1, column=1, value="Range Deleted")
    worksheet.cell(row=1, column=2, value="Cell Count")

    for row_index, range_text in enumerate(ranges_to_delete, start=2):
        worksheet.cell(row=row_index, column=1, value=range_text)
        if ":" not in range_text:
            cell_count = 1
        else:
            min_column, min_row, max_column, max_row = range_boundaries(range_text)
            cell_count = (max_row - min_row + 1) * (max_column - min_column + 1)
        worksheet.cell(row=row_index, column=2, value=cell_count)

    summary_row = len(ranges_to_delete) + 3
    worksheet.cell(row=summary_row, column=1, value="Total cells deleted:")
    worksheet.cell(row=summary_row, column=2, value=deleted_cell_count)


def create_sheet_with_ranges_deleted(
    file_path: str | Path,
    sheet_name: str,
    ranges_to_delete: list[str],
    output_dir: str | Path,
    *,
    include_summary_sheet: bool = True,
) -> Path:
    output_path = output_workbook_path(file_path, sheet_name, "ranges_deleted", output_dir)
    source_workbook = load_workbook(file_path, data_only=False)
    target_workbook = Workbook()

    try:
        if sheet_name not in source_workbook.sheetnames:
            raise ValueError(f"Sheet '{sheet_name}' not found. Available: {source_workbook.sheetnames}")

        source_worksheet = source_workbook[sheet_name]
        target_worksheet = target_workbook.active
        target_worksheet.title = sheet_name
        copy_default_font(source_workbook, target_workbook)

        actual_max_row = source_worksheet.max_row
        actual_max_column = source_worksheet.max_column
        cells_to_delete = cells_for_ranges(ranges_to_delete, actual_max_row, actual_max_column)

        for row_index in range(1, actual_max_row + 1):
            for column_index in range(1, actual_max_column + 1):
                if (row_index, column_index) in cells_to_delete:
                    continue

                source_cell = source_worksheet.cell(row=row_index, column=column_index)
                if source_cell.value is None and not source_cell.has_style:
                    continue

                target_cell = target_worksheet.cell(row=row_index, column=column_index)
                copy_cell(source_cell, target_cell)

        copy_dimensions(source_worksheet, target_worksheet)
        copy_other_sheets(source_workbook, target_workbook, skip_sheet=sheet_name)

        if include_summary_sheet:
            summary_sheet = target_workbook.create_sheet("Deletion_Summary")
            write_deletion_summary(summary_sheet, ranges_to_delete, len(cells_to_delete))

        set_automatic_calculation(target_workbook)
        target_workbook.save(output_path)
        return output_path
    finally:
        source_workbook.close()
        target_workbook.close()
