from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell import Cell
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.utils.cell import coordinate_to_tuple


CELL_REF_RE = re.compile(
    r"(?<![A-Za-z!])"
    r"\$?([A-Z]{1,3})"
    r"\$?(\d{1,7})"
    r"(?![A-Za-z\d(])"
)
FORMULA_REGION_TYPES = {"calculation_section", "data_table"}


def iter_excel_range(worksheet, range_text: str):
    result = worksheet[range_text]
    if isinstance(result, Cell):
        yield (result,)
    else:
        yield from result


def region_id(region: dict) -> int:
    value = region.get("id")
    if not isinstance(value, int):
        raise ValueError("Region id must be an integer")
    return value


def split_formula_regions(regions: list[dict]) -> tuple[list[dict], list[dict]]:
    formula_regions = []
    other_regions = []
    for region in regions:
        if region.get("type") in FORMULA_REGION_TYPES:
            formula_regions.append(region)
        else:
            other_regions.append(region)
    return formula_regions, other_regions


def region_contains_cell(worksheet, region: dict, active_cell: str) -> bool:
    region_range = region.get("range")
    if not isinstance(region_range, str) or not region_range:
        return False

    for row in iter_excel_range(worksheet, region_range):
        for cell in row:
            if cell.coordinate == active_cell:
                return True
    return False


def find_active_region(worksheet, regions: list[dict], active_cell: str) -> dict | None:
    for region in regions:
        if region_contains_cell(worksheet, region, active_cell):
            return region
    return None


def nearest_region_for_cell(regions: list[dict], active_cell: str) -> dict | None:
    active_row, active_column = coordinate_to_tuple(active_cell)
    best_region = None
    best_distance = float("inf")

    for region in regions:
        try:
            min_column, min_row, max_column, max_row = range_boundaries(region.get("range"))
        except Exception:
            continue

        row_distance = max(0, min_row - active_row, active_row - max_row)
        column_distance = max(0, min_column - active_column, active_column - max_column)
        distance = row_distance + column_distance
        if distance < best_distance:
            best_distance = distance
            best_region = region

    return best_region


def expand_region_to_include_cell(region: dict, active_cell: str) -> dict:
    active_row, active_column = coordinate_to_tuple(active_cell)
    min_column, min_row, max_column, max_row = range_boundaries(region["range"])

    expanded_range = (
        f"{get_column_letter(min(min_column, active_column))}{min(min_row, active_row)}"
        f":{get_column_letter(max(max_column, active_column))}{max(max_row, active_row)}"
    )
    return {**region, "range": expanded_range}


def regions_by_id(regions: list[dict]) -> dict[int, dict]:
    return {region_id(region): region for region in regions}


def add_dependency_regions(
    focused_regions: dict[int, dict],
    selected_region: dict,
    all_regions_by_id: dict[int, dict],
) -> None:
    for dependency_id in selected_region.get("dependencies", []) or []:
        if dependency_id not in all_regions_by_id:
            raise ValueError(f"Dependency region {dependency_id} was not found")
        focused_regions[dependency_id] = all_regions_by_id[dependency_id]


def add_preceding_regions(
    focused_regions: dict[int, dict],
    selected_region: dict,
    regions: list[dict],
) -> None:
    selected_region_id = region_id(selected_region)
    for region in regions:
        current_region_id = region_id(region)
        if current_region_id < selected_region_id:
            focused_regions[current_region_id] = region


def expand_regions_for_headers(
    filepath: str | Path,
    sheet_name: str,
    focused_regions: dict[int, dict],
    *,
    max_header_rows: int = 3,
) -> dict[int, dict]:
    workbook = load_workbook(filepath, data_only=False)
    try:
        worksheet = workbook[sheet_name]
        expanded_regions = {}

        for current_region_id, region in focused_regions.items():
            try:
                min_column, min_row, max_column, max_row = range_boundaries(region.get("range"))
            except Exception:
                expanded_regions[current_region_id] = region
                continue

            new_min_row = min_row
            consecutive_empty_rows = 0
            for offset in range(1, max_header_rows + 1):
                row_index = min_row - offset
                if row_index < 1:
                    break

                has_content = any(
                    worksheet.cell(row_index, column_index).value is not None
                    for column_index in range(min_column, max_column + 1)
                )
                if has_content:
                    new_min_row = row_index
                    consecutive_empty_rows = 0
                else:
                    consecutive_empty_rows += 1
                    if consecutive_empty_rows > 1:
                        break

            if new_min_row == min_row:
                expanded_regions[current_region_id] = region
                continue

            expanded_range = f"{get_column_letter(min_column)}{new_min_row}:{get_column_letter(max_column)}{max_row}"
            expanded_regions[current_region_id] = {**region, "range": expanded_range}

        return expanded_regions
    finally:
        workbook.close()


def formula_referenced_rows(worksheet, focused_regions: dict[int, dict]) -> set[int]:
    referenced_rows = set()
    for region in focused_regions.values():
        try:
            min_column, min_row, max_column, max_row = range_boundaries(region.get("range"))
        except Exception:
            continue

        for row_index in range(min_row, max_row + 1):
            for column_index in range(min_column, max_column + 1):
                value = worksheet.cell(row_index, column_index).value
                if value and str(value).startswith("="):
                    for _, row_text in CELL_REF_RE.findall(str(value)):
                        referenced_rows.add(int(row_text))
    return referenced_rows


def row_to_unfocused_region(
    all_regions: list[dict],
    focused_region_ids: set[int],
) -> dict[int, dict]:
    row_mapping = {}
    for region in all_regions:
        current_region_id = region_id(region)
        if current_region_id in focused_region_ids:
            continue

        try:
            _, min_row, _, max_row = range_boundaries(region.get("range"))
        except Exception:
            continue

        for row_index in range(min_row, max_row + 1):
            row_mapping[row_index] = region
    return row_mapping


def include_formula_reference_regions(
    filepath: str | Path,
    sheet_name: str,
    focused_regions: dict[int, dict],
    all_regions: list[dict],
) -> dict[int, dict]:
    workbook = load_workbook(filepath, data_only=False)
    try:
        worksheet = workbook[sheet_name]
        changed = True
        while changed:
            changed = False
            row_mapping = row_to_unfocused_region(all_regions, set(focused_regions.keys()))
            if not row_mapping:
                break

            for row_index in formula_referenced_rows(worksheet, focused_regions):
                if row_index not in row_mapping:
                    continue
                new_region = row_mapping[row_index]
                new_region_id = region_id(new_region)
                if new_region_id not in focused_regions:
                    focused_regions[new_region_id] = new_region
                    changed = True
        return focused_regions
    finally:
        workbook.close()


def focused_row_ranges(focused_regions: dict[int, dict]) -> list[tuple[int, int, int]]:
    row_ranges = []
    for current_region_id, region in focused_regions.items():
        try:
            _, min_row, _, max_row = range_boundaries(region.get("range"))
        except Exception:
            continue
        row_ranges.append((min_row, max_row, current_region_id))
    return row_ranges


def referenced_gap_rows(worksheet, focused_regions: dict[int, dict]) -> set[int]:
    row_ranges = focused_row_ranges(focused_regions)
    covered_rows = set()
    for min_row, max_row, _ in row_ranges:
        covered_rows.update(range(min_row, max_row + 1))

    needed_rows = set()
    for row_index in formula_referenced_rows(worksheet, focused_regions):
        if row_index < 1 or row_index > 1048576:
            continue
        if row_index in covered_rows:
            continue
        if any(worksheet.cell(row_index, column_index).value is not None for column_index in range(1, worksheet.max_column + 1)):
            needed_rows.add(row_index)
    return needed_rows


def expand_nearest_region_for_gap_rows(
    focused_regions: dict[int, dict],
    needed_rows: set[int],
) -> dict[int, dict]:
    row_ranges = focused_row_ranges(focused_regions)
    for row_index in needed_rows:
        best_region_id = None
        best_distance = float("inf")

        for min_row, max_row, current_region_id in row_ranges:
            distance = min(abs(row_index - min_row), abs(row_index - max_row))
            if distance < best_distance:
                best_distance = distance
                best_region_id = current_region_id

        if best_region_id is None:
            continue

        region = focused_regions[best_region_id]
        min_column, min_row, max_column, max_row = range_boundaries(region.get("range"))
        expanded_range = (
            f"{get_column_letter(min_column)}{min(min_row, row_index)}"
            f":{get_column_letter(max_column)}{max(max_row, row_index)}"
        )
        focused_regions[best_region_id] = {**region, "range": expanded_range}

    return focused_regions


def close_formula_references(
    filepath: str | Path,
    sheet_name: str,
    focused_regions: dict[int, dict],
    all_regions: list[dict],
) -> dict[int, dict]:
    focused_regions = include_formula_reference_regions(filepath, sheet_name, focused_regions, all_regions)

    workbook = load_workbook(filepath, data_only=False)
    try:
        worksheet = workbook[sheet_name]
        needed_rows = referenced_gap_rows(worksheet, focused_regions)
    finally:
        workbook.close()

    if needed_rows:
        focused_regions = expand_nearest_region_for_gap_rows(focused_regions, needed_rows)
    return focused_regions


def focused_regions_for_active_cell(
    filepath: str | Path,
    sheet_name: str,
    active_cell: str,
    regions: list[dict],
) -> dict[int, dict]:
    if not regions:
        raise ValueError("No regions found in segmentation data")

    active_cell = active_cell.upper()
    formula_regions, _ = split_formula_regions(regions)

    workbook = load_workbook(filepath, data_only=False)
    try:
        worksheet = workbook[sheet_name]
        selected_region = find_active_region(worksheet, formula_regions, active_cell)
        if selected_region is None:
            selected_region = find_active_region(worksheet, regions, active_cell)
        if selected_region is None:
            selected_region = nearest_region_for_cell(regions, active_cell)
            if selected_region is not None:
                selected_region = expand_region_to_include_cell(selected_region, active_cell)
    finally:
        workbook.close()

    if selected_region is None:
        raise ValueError(f"Active cell {active_cell} not found in any region")

    selected_region_id = region_id(selected_region)
    focused_regions = {selected_region_id: selected_region}
    all_regions_by_id = regions_by_id(regions)

    add_dependency_regions(focused_regions, selected_region, all_regions_by_id)
    add_preceding_regions(focused_regions, selected_region, regions)
    focused_regions = expand_regions_for_headers(filepath, sheet_name, focused_regions)
    focused_regions = close_formula_references(filepath, sheet_name, focused_regions, regions)
    return focused_regions
