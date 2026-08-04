"""Formula reference extraction and range-overlap utilities."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.utils.cell import column_index_from_string, coordinate_to_tuple, get_column_letter


EXCEL_MAX_ROWS = 1_048_576
EXCEL_MAX_COLUMNS = 16_384
DEFAULT_MAX_EXPANDED_CELLS = 250_000

_CELL_RE = re.compile(r"^\$?[A-Z]{1,3}\$?\d+#?$")
_COLUMN_RE = re.compile(r"^\$?[A-Z]{1,3}#?$")
_ROW_RE = re.compile(r"^\$?\d+#?$")
_INVALID_AT_REF_RE = re.compile(r"\$?[A-Z]{1,3}@")
_FUNC_REF_PREFIX_RE = re.compile(r"(?P<ref>(?:'[^']+'!|[A-Za-z_][\w.]*!)?\$?[A-Z]{1,3}\$?\d+#?)\s*:")
_FALLBACK_REF_RE = re.compile(
    r"(?:(?P<sheet>'(?:[^']|'')*'|[A-Za-z_][\w.]*)!)?"
    r"(?P<ref>"
    r"\$?[A-Z]{1,3}\$?\d+#?(?::\$?[A-Z]{1,3}\$?\d+#?)?"
    r"|\$?[A-Z]{1,3}#?:\$?[A-Z]{1,3}#?"
    r"|\$?\d+#?:\$?\d+#?"
    r")"
)


class TokenizeError(ValueError):
    pass


@dataclass
class FormulaReferenceSet:
    """Expanded referenced cells plus unresolved reference diagnostics."""

    cells: set[str] = field(default_factory=set)
    unresolved: set[str] = field(default_factory=set)
    warnings: list[str] = field(default_factory=list)

    def metadata(self) -> dict:
        return {
            "cell_count": len(self.cells),
            "unresolved": sorted(self.unresolved),
            "warnings": list(self.warnings),
        }


@dataclass
class RangeOverlapResult:
    """Expected-reference recall for a predicted formula."""

    value: float
    overlap_count: int
    expected_count: int
    predicted_count: int
    expected_unresolved: set[str] = field(default_factory=set)
    predicted_unresolved: set[str] = field(default_factory=set)
    warnings: list[str] = field(default_factory=list)

    def metadata(self) -> dict:
        return {
            "value": self.value,
            "overlap_count": self.overlap_count,
            "expected_count": self.expected_count,
            "predicted_count": self.predicted_count,
            "expected_unresolved": sorted(self.expected_unresolved),
            "predicted_unresolved": sorted(self.predicted_unresolved),
            "warnings": list(self.warnings),
        }


def extract_referenced_cells(
    formula: str | None,
    *,
    workbook_path: Path | str | None = None,
    worksheet_name: str | None = None,
    active_cell: str | None = None,
    max_expanded_cells: int = DEFAULT_MAX_EXPANDED_CELLS,
) -> FormulaReferenceSet:
    """Return a normalized, sheet-qualified set of cells referenced by ``formula``."""

    extractor = _ReferenceExtractor(
        workbook_path=workbook_path,
        worksheet_name=worksheet_name,
        active_cell=active_cell,
        max_expanded_cells=max_expanded_cells,
    )
    try:
        return extractor.extract(formula)
    finally:
        extractor.close()


def range_overlap(
    predicted_formula: str | None,
    expected_formula: str | None,
    *,
    workbook_path: Path | str | None = None,
    worksheet_name: str | None = None,
    active_cell: str | None = None,
    max_expanded_cells: int = DEFAULT_MAX_EXPANDED_CELLS,
) -> RangeOverlapResult:
    """Measure expected-cell coverage by the predicted formula."""

    extractor = _ReferenceExtractor(
        workbook_path=workbook_path,
        worksheet_name=worksheet_name,
        active_cell=active_cell,
        max_expanded_cells=max_expanded_cells,
    )
    try:
        expected = extractor.extract(expected_formula)
        predicted = extractor.extract(predicted_formula)
    finally:
        extractor.close()

    overlap = expected.cells & predicted.cells
    if expected.cells:
        value = len(overlap) / len(expected.cells)
    else:
        value = 0.0 if expected.unresolved else 1.0

    return RangeOverlapResult(
        value=value,
        overlap_count=len(overlap),
        expected_count=len(expected.cells),
        predicted_count=len(predicted.cells),
        expected_unresolved=set(expected.unresolved),
        predicted_unresolved=set(predicted.unresolved),
        warnings=[*expected.warnings, *predicted.warnings],
    )


class _ReferenceExtractor:
    def __init__(
        self,
        *,
        workbook_path: Path | str | None,
        worksheet_name: str | None,
        active_cell: str | None,
        max_expanded_cells: int,
    ) -> None:
        self.workbook_path = Path(workbook_path) if workbook_path is not None else None
        self.worksheet_name = worksheet_name
        self.active_cell = _normalize_cell(active_cell) if active_cell else None
        self.max_expanded_cells = max_expanded_cells
        self._workbook = None

    def close(self) -> None:
        if self._workbook is not None:
            self._workbook.close()
            self._workbook = None

    def extract(self, formula: str | None) -> FormulaReferenceSet:
        result = FormulaReferenceSet()
        normalized = _normalize_formula(formula)
        if not normalized:
            return result

        try:
            self._extract_with_tokenizer(normalized, result)
        except Exception as exc:
            result.warnings.append(f"parser_fallback:{exc.__class__.__name__}")
            self._extract_with_regex(normalized, result)
        return result

    def _extract_with_tokenizer(self, formula: str, result: FormulaReferenceSet) -> None:
        if _INVALID_AT_REF_RE.search(formula):
            raise TokenizeError()
        tokens = Tokenizer(formula).items
        local_names = _let_local_names(tokens)
        for token in tokens:
            if token.type == "FUNC" and token.subtype == "OPEN":
                self._add_function_reference_prefix(token.value, result)
                continue
            if token.type != "OPERAND" or token.subtype != "RANGE":
                continue
            self._add_reference_token(token.value, local_names, result)

    def _add_function_reference_prefix(self, value: str, result: FormulaReferenceSet) -> None:
        match = _FUNC_REF_PREFIX_RE.match(value)
        if not match:
            return
        self._add_reference_token(match.group("ref"), set(), result)

    def _add_reference_token(self, value: str, local_names: set[str], result: FormulaReferenceSet) -> None:
        token = value.strip()
        if not token:
            return
        book, sheet, ref = _split_reference_token(token)
        if book:
            result.unresolved.add(f"[{book}]")
            return
        if _is_structured_reference(ref):
            self._add_structured_reference(ref, sheet, result)
            return
        clean_ref = _normalize_ref_token(ref)
        if clean_ref.startswith(":"):
            return
        if clean_ref.upper() in local_names:
            return
        if ":" in clean_ref:
            left, right = clean_ref.split(":", 1)
            self._add_range(left, right, sheet, book, result)
            return
        if _CELL_RE.match(clean_ref):
            self._add_cell(clean_ref, sheet, book, result)
            return
        if _COLUMN_RE.match(clean_ref) or _ROW_RE.match(clean_ref):
            result.unresolved.add(_format_ref(sheet, ref, book))
            return
        self._add_named_range(ref, book, result)

    def _add_cell(
        self,
        cell: str,
        sheet: str | None,
        book: str | None,
        result: FormulaReferenceSet,
    ) -> None:
        try:
            normalized = _normalize_cell(cell)
            row, col = coordinate_to_tuple(normalized)
        except Exception:
            result.unresolved.add(_format_ref(sheet, cell, book))
            return
        result.cells.add(self._cell_key(sheet, row, col, book))

    def _add_range(
        self,
        left: str,
        right: str,
        sheet: str | None,
        book: str | None,
        result: FormulaReferenceSet,
    ) -> None:
        left = _normalize_ref_token(left)
        right = _normalize_ref_token(right)
        sheet_name = self._effective_sheet(sheet)
        label = _format_ref(sheet_name, f"{left}:{right}", book)

        if _CELL_RE.match(left) and _CELL_RE.match(right):
            r1, c1 = coordinate_to_tuple(left)
            r2, c2 = coordinate_to_tuple(right)
            self._expand_bounds(sheet, min(r1, r2), max(r1, r2), min(c1, c2), max(c1, c2), book, label, result)
            return

        bounds = self._sheet_bounds(sheet_name, result)
        if bounds is None:
            result.unresolved.add(label)
            return

        max_row, max_col = bounds
        if _COLUMN_RE.match(left) and _COLUMN_RE.match(right):
            c1 = column_index_from_string(left)
            c2 = column_index_from_string(right)
            self._expand_bounds(sheet, 1, max_row, min(c1, c2), max(c1, c2), book, label, result)
            return

        if _ROW_RE.match(left) and _ROW_RE.match(right):
            r1 = int(left)
            r2 = int(right)
            self._expand_bounds(sheet, min(r1, r2), max(r1, r2), 1, max_col, book, label, result)
            return

        result.unresolved.add(label)

    def _expand_bounds(
        self,
        sheet: str | None,
        min_row: int,
        max_row: int,
        min_col: int,
        max_col: int,
        book: str | None,
        label: str,
        result: FormulaReferenceSet,
    ) -> None:
        total = (max_row - min_row + 1) * (max_col - min_col + 1)
        if total > self.max_expanded_cells:
            result.unresolved.add(label)
            result.warnings.append(f"range_too_large:{label}:{total}")
            return
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                result.cells.add(self._cell_key(sheet, row, col, book))

    def _add_named_range(self, name: str, book: str | None, result: FormulaReferenceSet) -> None:
        workbook = self._load_workbook(result)
        if workbook is None:
            result.unresolved.add(name)
            return

        definitions = _find_defined_names(workbook, name)
        if not definitions:
            result.unresolved.add(name)
            return

        for sheet, coord in definitions:
            self._add_defined_name_destination(sheet, coord, book, result)

    def _add_defined_name_destination(
        self,
        sheet: str,
        coord: str,
        book: str | None,
        result: FormulaReferenceSet,
    ) -> None:
        coord = coord.replace("$", "")
        for part in coord.split(","):
            ref = part.strip()
            if not ref:
                continue
            if "!" in ref:
                sheet_part, ref = ref.rsplit("!", 1)
                sheet = _normalize_sheet(sheet_part)
            if ":" in ref:
                left, right = ref.split(":", 1)
                self._add_range(left, right, sheet, book, result)
            else:
                self._add_cell(ref, sheet, book, result)

    def _add_structured_reference(self, token: str, sheet: str | None, result: FormulaReferenceSet) -> None:
        table_name, columns, items, this_row = _parse_structured_reference(token)
        if not table_name:
            result.unresolved.add(token)
            return
        table_info = self._find_table(table_name, result)
        if table_info is None:
            result.unresolved.add(token)
            return
        sheet_name, table = table_info
        table_columns = _table_column_names(table)
        if not columns:
            col_indexes = range(1, len(table_columns) + 1)
        else:
            try:
                indexes = [table_columns.index(column) + 1 for column in columns]
            except ValueError:
                result.unresolved.add(token)
                return
            if len(indexes) == 2 and ":" in token:
                col_indexes = range(min(indexes), max(indexes) + 1)
            else:
                col_indexes = indexes
        self._expand_table_columns(sheet_name if sheet is None else sheet, table, col_indexes, items, this_row, None, token, result)

    def _expand_table_columns(
        self,
        sheet: str,
        table,
        col_indexes: Iterable[int],
        items: list[str],
        this_row: bool,
        book: str | None,
        label: str,
        result: FormulaReferenceSet,
    ) -> None:
        min_col, min_row, max_col, max_row = _range_boundaries(table.ref)
        rows = self._table_rows(min_row, max_row, table, items, this_row, label, result)
        if rows is None:
            return
        for index in col_indexes:
            col = min_col + index - 1
            self._expand_bounds(sheet, rows[0], rows[1], col, col, book, label, result)

    def _table_rows(
        self,
        min_row: int,
        max_row: int,
        table,
        items: list[str],
        this_row: bool,
        label: str,
        result: FormulaReferenceSet,
    ) -> tuple[int, int] | None:
        if this_row or "This Row" in items:
            if not self.active_cell:
                result.unresolved.add(label)
                return None
            row, _ = coordinate_to_tuple(self.active_cell)
            return row, row

        header_rows = int(getattr(table, "headerRowCount", 1) or 0)
        totals_rows = int(getattr(table, "totalsRowCount", 0) or 0)
        data_start = min_row + header_rows
        data_end = max_row - totals_rows

        if "All" in items:
            return min_row, max_row
        if "Headers" in items:
            return min_row, min_row + max(header_rows - 1, 0)
        if "Totals" in items:
            return max_row - max(totals_rows - 1, 0), max_row
        return data_start, data_end

    def _find_table(self, table_name: str, result: FormulaReferenceSet):
        workbook = self._load_workbook(result)
        if workbook is None:
            return None
        for worksheet in workbook.worksheets:
            if table_name in worksheet.tables:
                return worksheet.title, worksheet.tables[table_name]
        return None

    def _sheet_bounds(self, sheet: str | None, result: FormulaReferenceSet) -> tuple[int, int] | None:
        workbook = self._load_workbook(result)
        if workbook is None:
            return None
        sheet_name = self._effective_sheet(sheet)
        if sheet_name not in workbook.sheetnames:
            result.warnings.append(f"sheet_not_found:{sheet_name}")
            return None
        worksheet = workbook[sheet_name]
        return worksheet.max_row or EXCEL_MAX_ROWS, worksheet.max_column or EXCEL_MAX_COLUMNS

    def _load_workbook(self, result: FormulaReferenceSet):
        if self._workbook is not None:
            return self._workbook
        if self.workbook_path is None:
            return None
        try:
            self._workbook = load_workbook(self.workbook_path, read_only=False, data_only=False)
        except Exception as exc:
            result.warnings.append(f"workbook_load_failed:{exc.__class__.__name__}")
            return None
        return self._workbook

    def _cell_key(self, sheet: str | None, row: int, col: int, book: str | None) -> str:
        sheet_name = self._effective_sheet(sheet)
        cell = f"{get_column_letter(col)}{row}"
        book_prefix = f"[{book}]" if book else ""
        return f"{book_prefix}{sheet_name}!{cell}" if sheet_name else f"{book_prefix}{cell}"

    def _effective_sheet(self, sheet: str | None) -> str | None:
        return _normalize_sheet(sheet) if sheet else self.worksheet_name

    def _extract_with_regex(self, formula: str, result: FormulaReferenceSet) -> None:
        text = _strip_string_literals(formula)
        for match in _FALLBACK_REF_RE.finditer(text):
            sheet = _normalize_sheet(match.group("sheet"))
            ref = match.group("ref")
            if ":" in ref:
                left, right = ref.split(":", 1)
                self._add_range(left, right, sheet, None, result)
            else:
                self._add_cell(ref, sheet, None, result)


def _let_local_names(tokens) -> set[str]:
    local_names: set[str] = set()
    for index, token in enumerate(tokens):
        if token.type != "FUNC" or token.subtype != "OPEN" or token.value.upper() != "LET(":
            continue
        args = _function_arguments(tokens, index)
        for arg in args[:-1:2]:
            if len(arg) != 1:
                continue
            name_token = arg[0]
            if name_token.type == "OPERAND" and name_token.subtype == "RANGE" and _is_let_name(name_token.value):
                local_names.add(name_token.value.upper())
    return local_names


def _function_arguments(tokens, open_index: int) -> list[list]:
    args: list[list] = []
    current: list = []
    depth = 0
    for token in tokens[open_index + 1:]:
        if token.type == "FUNC" and token.subtype == "OPEN":
            depth += 1
            current.append(token)
            continue
        if token.type == "FUNC" and token.subtype == "CLOSE":
            if depth == 0:
                args.append(current)
                break
            depth -= 1
            current.append(token)
            continue
        if depth == 0 and token.type == "SEP" and token.subtype == "ARG":
            args.append(current)
            current = []
            continue
        current.append(token)
    return args


def _split_reference_token(token: str) -> tuple[str | None, str | None, str]:
    book = None
    sheet = None
    ref = token
    if "!" not in ref:
        return book, sheet, ref
    sheet_part, ref = ref.rsplit("!", 1)
    if sheet_part.startswith("[") and "]" in sheet_part:
        book, sheet_part = sheet_part[1:].split("]", 1)
    sheet = _normalize_sheet(sheet_part)
    return book, sheet, ref


def _is_bare_name(value: str) -> bool:
    token = value.strip()
    if not token or "!" in token or "[" in token or "]" in token or ":" in token:
        return False
    normalized = _normalize_ref_token(token)
    return not (_CELL_RE.match(normalized) or _COLUMN_RE.match(normalized) or _ROW_RE.match(normalized))


def _is_let_name(value: str) -> bool:
    token = value.strip()
    if not token or "!" in token or "[" in token or "]" in token or ":" in token:
        return False
    normalized = _normalize_ref_token(token)
    return not (_CELL_RE.match(normalized) or _ROW_RE.match(normalized))


def _is_structured_reference(ref: str) -> bool:
    return "[" in ref and "]" in ref


def _parse_structured_reference(ref: str) -> tuple[str | None, list[str], list[str], bool]:
    table_name, _, body = ref.partition("[")
    if not table_name:
        return None, [], [], False
    body = "[" + body
    parts = re.findall(r"\[([^\]]+)\]", body)
    if not parts:
        parts = [body.strip("[]")]

    items = []
    columns = []
    this_row = False
    for part in parts:
        cleaned = part.strip()
        if not cleaned:
            continue
        if cleaned.startswith("#"):
            label = cleaned[1:]
            if label == "All":
                items.append("All")
            elif label == "Headers":
                items.append("Headers")
            elif label == "Totals":
                items.append("Totals")
            elif label == "This Row":
                this_row = True
            continue
        if cleaned.startswith("@"):
            this_row = True
            cleaned = cleaned[1:]
        if cleaned and cleaned not in columns:
            columns.append(cleaned)
    return table_name, columns, items, this_row


def _normalize_formula(formula: str | None) -> str:
    if formula is None:
        return ""
    normalized = str(formula).strip()
    if normalized.startswith("{=") and normalized.endswith("}"):
        normalized = "=" + normalized[2:-1].strip()
    return (
        normalized.replace("_xlfn.", "")
        .replace("_XLFN.", "")
        .replace("_xlws.", "")
        .replace("_XLWS.", "")
    )


def _normalize_ref_token(ref: str) -> str:
    return ref.replace("$", "").replace("#", "").upper()


def _normalize_cell(cell: str) -> str:
    normalized = _normalize_ref_token(cell)
    row, col = coordinate_to_tuple(normalized)
    return f"{get_column_letter(col)}{row}"


def _normalize_sheet(sheet: str | None) -> str | None:
    if sheet is None:
        return None
    normalized = str(sheet).strip()
    if normalized.endswith("!"):
        normalized = normalized[:-1]
    if normalized.startswith("'") and normalized.endswith("'"):
        normalized = normalized[1:-1].replace("''", "'")
    return normalized


def _format_ref(sheet: str | None, ref: str, book: str | None) -> str:
    book_prefix = f"[{book}]" if book else ""
    sheet_prefix = f"{_normalize_sheet(sheet)}!" if sheet else ""
    return f"{book_prefix}{sheet_prefix}{ref}"


def _strip_string_literals(formula: str) -> str:
    return re.sub(r'"(?:[^"]|"")*"', '""', formula)


def _find_defined_names(workbook, name: str) -> list[tuple[str, str]]:
    definitions = []
    try:
        defined_name = workbook.defined_names[name]
        definitions.extend(list(defined_name.destinations))
    except Exception:
        pass

    try:
        iterable = workbook.defined_names.definedName
    except AttributeError:
        iterable = []

    for defined_name in iterable:
        if getattr(defined_name, "name", None) != name:
            continue
        try:
            definitions.extend(list(defined_name.destinations))
        except Exception:
            continue

    return definitions


def _table_column_names(table) -> list[str]:
    return [column.name for column in table.tableColumns]


def _range_boundaries(ref: str) -> tuple[int, int, int, int]:
    from openpyxl.utils.cell import range_boundaries

    return range_boundaries(ref)
